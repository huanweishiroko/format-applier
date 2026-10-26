# -*- coding: utf-8 -*-
"""首页：拖拽区 + 文件信息 / 参数双栏 + 操作按钮 + 可折叠日志。

Phase 3 已接通真实处理流程：
  * 预设下拉合并「内置 presets/」与「用户预设目录」，用 📦/👤 区分来源；
  * [预览效果] → 真跑一次处理到临时目录，弹窗展示标题结构 + 角色映射 + 前 3 段效果；
  * [开始处理] → bridge.process 输出到用户选定目录（后台线程 + 进度 + 逐行日志）；
  * [上传格式要求] → upload_dialog（txt/docx/自然语言 → DSL → 可存为用户预设）。
"""

from __future__ import annotations

import shutil
import sys
import tempfile
import threading
import time
import tkinter as tk
from pathlib import Path
from tkinter import filedialog
from typing import List, Optional

import customtkinter as ctk

# 绝对优先，理由见 preview.py 顶部注释（frozen 下 ui/core 是顶层包）
try:
    from ui import dialogs as dlg
    from ui import preview as pv
    from ui import theme as th
    from ui import upload_dialog as up_dlg
    from core import bridge
    from core import config as cfg
    from core import preset_loader as pl
    import dnd  # 顶层模块：Windows 原生 WM_DROPFILES 拖放
except ImportError:  # pragma: no cover - 包内运行时
    from . import dialogs as dlg  # type: ignore
    from . import preview as pv  # type: ignore
    from . import theme as th  # type: ignore
    from . import upload_dialog as up_dlg  # type: ignore
    from ..core import bridge  # type: ignore
    from ..core import config as cfg  # type: ignore
    from ..core import preset_loader as pl  # type: ignore
    from .. import dnd  # type: ignore


DOCX_SUFFIX = ".docx"

# 拖拽区四态（阶段三）：idle 常态 / hover 悬停 / pulse 接收脉冲 / selected 已选文件。
# 配色与线型一律走 theme 的 DROP_* token（本模块不再自带色常量），
# 高度由 _left 的 row1 minsize 控制（见 __init__，阶段三 184 → 200）。
DROP_HINT_IDLE = "支持单个文件；也可以直接点下面的按钮选择"
DROP_HINT_ACTIVE = "松开鼠标即可导入该文件"
DROP_HINT_SELECTED_FMT = "已选择：{}　拖入新文件可替换"
DROP_TITLE_IDLE = "拖拽 .docx 文件到此处"
DROP_TITLE_ACTIVE = "松开即可导入"

# 状态 → 渲染参数：(边框色对, 线宽, 虚线 gap, 底/填充色对, 图标, 标题, 提示)
# gap=1 是实线（_render_drop_border 的特殊分支），gap=7/4/8 才是虚线。
# 四态必须两两可辨：pulse 用 accent 淡底 + 绿描边（"正在接收"），
# selected 用绿淡底 + 绿实线（"已经选好"）——两者底色不同，不会混。
def _make_drop_states():
    """按当前（唯一）皮肤生成状态表。

    单皮肤后只在模块导入/页面构建时算一次；不再有运行时换肤。
    """
    states = {
        "idle": (th.DROP_BORDER_IDLE, 1, 7, th.DROP_FILL_IDLE,
                 "⬇", DROP_TITLE_IDLE, DROP_HINT_IDLE),
        "hover": (th.DROP_BORDER_ACTIVE, 3, 4, th.DROP_FILL_ACTIVE,
                  "📥", DROP_TITLE_ACTIVE, DROP_HINT_ACTIVE),
        "pulse": (th.DROP_BORDER_OK, 2, 8, th.DROP_FILL_ACTIVE,
                  "📥", DROP_TITLE_ACTIVE, DROP_HINT_ACTIVE),
        "selected": (th.DROP_BORDER_OK, 1.5, 1, th.DROP_FILL_OK,
                     "✅", DROP_TITLE_IDLE, None),
    }
    pulses = {
        "pulse_off": (th.DROP_BORDER_OK, 3, 8, th.DROP_FILL_ACTIVE,
                      "📥", DROP_TITLE_ACTIVE, DROP_HINT_ACTIVE),
        "pulse_on": (th.DROP_BORDER_OK, 1, 8, th.DROP_FILL_ACTIVE,
                     "📥", DROP_TITLE_ACTIVE, DROP_HINT_ACTIVE),
    }
    token_colors = tuple(
        value
        for value in (th.DROP_BORDER_IDLE, th.DROP_BORDER_ACTIVE,
                      th.DROP_FILL_IDLE, th.DROP_FILL_ACTIVE,
                      th.DROP_FILL_OK, th.DROP_BORDER_OK)
        if isinstance(value, str)
    )
    return states, pulses, token_colors


_DROP_STATES, _DROP_PULSE_SPECS, _DROP_TOKEN_COLORS = _make_drop_states()

# 120ms 状态过渡：4 步 × 30ms（只用 after()，不引新依赖）
_DROP_TWEEN_MS = 30
_DROP_TWEEN_STEPS = 4


def _hex_to_rgb(value: str) -> Optional[tuple]:
    """把 ``#RRGGBB`` 解成 (r, g, b)，解析不了返回 None。"""
    text = str(value).strip()
    if not text.startswith("#"):
        return None
    digits = text[1:]
    if len(digits) == 3:
        digits = "".join(ch * 2 for ch in digits)
    if len(digits) != 6:
        return None
    try:
        return (int(digits[0:2], 16), int(digits[2:4], 16), int(digits[4:6], 16))
    except ValueError:
        return None


def _blend_color(src: str, dst: str, mix: float) -> str:
    """在两个十六进制色之间线性插值；任一端解不出来就退化为 dst。"""
    if mix <= 0:
        return str(src)
    if mix >= 1:
        return str(dst)
    a = _hex_to_rgb(src)
    b = _hex_to_rgb(dst)
    if a is None or b is None:
        return str(dst)
    return "#{:02x}{:02x}{:02x}".format(
        *(round(a[i] + (b[i] - a[i]) * mix) for i in range(3)))


def _blend_spec(src, dst, mix: float) -> tuple:
    """把 ``(浅, 深)`` 色对逐通道插值，保证过渡期间主题仍然成对可用。"""
    if isinstance(src, (tuple, list)) and isinstance(dst, (tuple, list)) \
            and len(src) == 2 and len(dst) == 2:
        return (_blend_color(src[0], dst[0], mix), _blend_color(src[1], dst[1], mix))
    if isinstance(src, (tuple, list)) and len(src) == 2:
        return src
    return dst


class HomePage(ctk.CTkFrame):
    """主页面。所有跨页状态都从 self.app 取，页面自己不持有全局配置。"""

    # 步骤轨道文案（Step 3 阶段二）
    STEP_TEXTS = ("拖入文件", "选择格式", "开始处理")

    def __init__(self, master, app):
        # 不透明页底（原来是 "transparent"）：4 页同格叠放 + tkraise 切换后，
        # 透明页之间没有遮挡层，下层页面的像素会从上层未绘制处透出来 → 观感是
        # "两页混合几帧"。填成 BG_PAGE 后每页自遮，视觉与原来一致（父容器底同色）。
        super().__init__(master, corner_radius=0, fg_color=th.BG_PAGE)
        self.app = app
        self.selected_file: Optional[Path] = None
        self._display_to_name: dict = {}
        self._busy = False
        self._worker_result: Optional[dict] = None
        self._worker_error: Optional[tuple] = None
        self._drag_active = False   # 鼠标是否正按住并停在窗口上（拖放悬停）
        # 拖拽区四态（阶段三）：idle / hover / pulse / selected。
        # flash 相关变量语义不变：_flash_steps 由 _take_paths 置位，_flash_step 常驻消费。
        self._flash_steps = 0
        self._flash_on = False
        self.after(220, self._flash_step)
        self._drop_state = "idle"       # 逻辑态（含 pulse）
        self._drop_dim = False          # True = 当前没有已选文件，不允许停在 selected
        self._drop_target = _DROP_STATES["idle"]   # 当前态的参数表
        self._drop_from = _DROP_STATES["idle"]     # 过渡起点（色对，用于插值）
        self._drop_mix = 1.0            # 过渡进度：1.0 = 已完全落到 _drop_target
        self._drop_step = 0             # 过渡已走步数
        self._drop_seq = 0              # 动画世代号：新一轮自增，旧帧据此自杀
        self._drop_bg_frame = th.C(th.DROP_FILL_IDLE)   # 最近一次真正刷进控件的底色（已解单色）
        self._drop_bg_solid = th.C(th.DROP_FILL_IDLE)   # 最近一次刷给 _canvas 的单色，兜底用
        self._drop_dash = (7, 1)        # (gap, line_width) 最近一次画的边框参数
        self._drop_line = th.C(th.DROP_BORDER_IDLE)     # 最近一次画的边框单色
        # 步骤轨道第 3 节点「开始处理」的纯视觉标记：上一轮是否成功处理完。
        # 不参与任何业务判定（是否可处理仍由 _require_ready 决定）。
        self._processed_ok = False

        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(0, weight=1)

        self._scroll = ctk.CTkScrollableFrame(self, corner_radius=0, fg_color="transparent")
        self._scroll.grid(row=0, column=0, sticky="nsew")
        self._scroll.grid_rowconfigure(0, weight=1)
        self._scroll.grid_columnconfigure(0, weight=1)      # 左列：唯一有 weight 的列，吸收余宽
        self._scroll.grid_columnconfigure(1, weight=0)      # 右列：宽度由 _right 的 320 决定

        # 右主左辅：左列 = 标题 + 拖拽区 + 已选文件卡 + 日志；右列 = 参数卡 + 按钮行。
        # 用「左 flex + 右固定 320」而非纯比例，参数卡宽度在任意窗口尺寸下都一样。
        self._left = ctk.CTkFrame(self._scroll, fg_color="transparent")
        self._left.grid(row=0, column=0, sticky="nsew", padx=(th.PAD_L, 0))
        self._left.grid_columnconfigure(0, weight=1)

        self._right = ctk.CTkFrame(self._scroll, fg_color="transparent", width=320)
        self._right.grid(row=0, column=1, sticky="nsew", padx=(th.PAD, th.PAD_L))
        self._right.grid_propagate(False)
        self._right.grid_columnconfigure(0, weight=1)

        # 左列行定义：0 标题 / 1 拖拽区 / 2 已选文件卡 / 3 日志。
        # 余高只认 row3（日志）——窗口加高时只有日志变长，拖拽区纹丝不动。
        # 拖拽区高度留给 row1 的 minsize（阶段三置为 200），行本身 weight=0。
        self._left.grid_rowconfigure(0, weight=0)
        self._left.grid_rowconfigure(1, weight=0, minsize=200)   # 拖拽区 200px（阶段三）
        self._left.grid_rowconfigure(2, weight=0)
        self._left.grid_rowconfigure(3, weight=1, minsize=160)

        self._build_header()
        self._build_drop_zone()
        self._build_columns()
        self._build_actions()
        self._build_log()

        self._stretch_scroll_viewport()
        self.refresh_presets()

    # ------------------------------------------------------------------ #
    # 1. 标题
    # ------------------------------------------------------------------ #

    def _build_header(self) -> None:
        head = ctk.CTkFrame(self._left, fg_color="transparent")
        head.grid(row=0, column=0, sticky="ew", pady=(th.PAD_L, th.PAD))
        head.grid_columnconfigure(0, weight=1)   # 标题列：吸收余宽
        head.grid_columnconfigure(1, weight=0)   # 轨道列：宽度由内容决定，贴右

        title_row = ctk.CTkFrame(head, fg_color="transparent")
        title_row.grid(row=0, column=0, sticky="w")
        title_font = th.font(28, "bold", role="display")
        ctk.CTkLabel(
            title_row, text=">", font=title_font,
            text_color=th.ACCENT, anchor="w",
        ).grid(row=0, column=0, sticky="w")
        ctk.CTkLabel(
            title_row, text="格式处理", font=title_font,
            text_color=th.TEXT, anchor="w",
        ).grid(row=0, column=1, sticky="w", padx=(th.SP.XS, 0))
        th.hint(
            head,
            "把 .docx 拖进来，选一个格式要求预设，剩下的交给桃符的格式处理引擎。",
        ).grid(row=1, column=0, sticky="w", pady=(4, 0))

        self._build_steps(head)

    # ------------------------------------------------------------------ #
    # 1a. 步骤轨道（拖入文件 → 选择格式 → 开始处理）
    # ------------------------------------------------------------------ #

    def _build_steps(self, master) -> None:
        """标题右侧的三节点轨道。只在这里画一次，之后一律由 :meth:`_refresh_steps` 改外观。

        颜色全部传 ``th`` 的 token 色对（不调 ``th.C()``）：PALETTE 每条都是 (深, 深)，
        CTk 自己解析；单皮肤固定深色，不存在运行时明暗切换。
        """
        steps = ctk.CTkFrame(master, fg_color="transparent")
        steps.grid(row=0, column=1, sticky="e", padx=(th.PAD, 0))

        self._step_dots: List[ctk.CTkLabel] = []
        self._step_labels: List[ctk.CTkLabel] = []
        self._step_lines: List[ctk.CTkFrame] = []

        for i, text in enumerate(self.STEP_TEXTS):
            dot = ctk.CTkLabel(
                steps, text="○", width=12,
                font=th.font(th.FS.SMALL), text_color=th.TEXT_MUTED,
            )
            dot.grid(row=0, column=i * 3, sticky="w")
            label = ctk.CTkLabel(
                steps, text=text,
                font=th.font(th.FS.SMALL), text_color=th.TEXT_MUTED,
            )
            label.grid(row=0, column=i * 3 + 1, sticky="w", padx=(th.SP.XS, 0))
            self._step_dots.append(dot)
            self._step_labels.append(label)
            if i < len(self.STEP_TEXTS) - 1:
                line = ctk.CTkFrame(
                    steps, width=18, height=1, corner_radius=0, fg_color=th.DIVIDER,
                )
                line.grid(row=0, column=i * 3 + 2, padx=th.SP.S)
                self._step_lines.append(line)

    def _refresh_steps(self) -> None:
        """按真实状态点亮轨道（累积式：第 N 个亮 = 前 N 个条件都满足）。

        只读三个已有事实，不新增业务判断：
        ① 有文件（``selected_file``）② 有预设（``current_preset()``）
        ③ 上一轮处理成功（``_processed_ok``，由 ``_on_process_done`` 回填）。

        累积式是为了避免「步骤 2 亮、步骤 1 灭」的错位观感：启动时 ``refresh_presets``
        会自动选中第一个预设，独立式下会一上来就点亮第 2 个节点。
        """
        if not getattr(self, "_step_dots", None):
            return
        raw = (
            self.selected_file is not None,
            getattr(self, "preset_var", None) is not None and self.current_preset() is not None,
            bool(self._processed_ok),
        )
        on = [all(raw[: i + 1]) for i in range(len(raw))]

        for i, (dot, label) in enumerate(zip(self._step_dots, self._step_labels)):
            lit = on[i]
            dot.configure(
                text="●" if lit else "○",
                text_color=th.ACCENT if lit else th.TEXT_MUTED,
            )
            label.configure(
                text_color=th.TEXT if lit else th.TEXT_MUTED,
                font=th.font(th.FS.SMALL, "bold" if lit else "normal"),
            )
        for i, line in enumerate(self._step_lines):
            line.configure(fg_color=th.ACCENT if on[i + 1] else th.DIVIDER)

    # ------------------------------------------------------------------ #
    # 1b. 让滚动容器内框吸收视口余高
    # ------------------------------------------------------------------ #

    def _stretch_scroll_viewport(self) -> None:
        """把「窗口比内容高」的那部分余高交给左列 row3（日志）吸收。

        customtkinter 6.0 的 ``CTkScrollableFrame._fit_frame_dimensions_to_canvas``
        （``ctk_scrollable_frame.py:250-254``）只把内框宽度对齐视口，**高度不设**，
        内框因此始终保持「内容自然高」。于是窗口加高时 ``_left`` 撑不到视口高度，
        row3 的 ``weight=1`` 无余高可吸（日志恒 150）。这里补上缺的那一半。

        做法是给内框的 canvas window item 设 height（``_set_dimensions`` 只改
        canvas 自身的 wanted height，动不了这个 item，实测无效）。
        内容超出视口时完全不抬升，滚动行为不受影响。

        常驻映射（2026-10-08 起 4 页同格叠放 + ``tkraise``）之后本方法有两个入口：
        ``<Configure>`` 与 :meth:`on_shown`。实测**叠在下层的 home 照样收到
        ``<Configure>``**（Tk 几何传播跟的是 map 状态，不是可见性），所以不需要
        "只在可见时重算"那种更严的 guard —— 用 ``winfo_ismapped()`` 即可，
        它既能挡住未映射时的误算，又不会误伤后台的合法重算。
        真正的重算逻辑在 :meth:`_fit_scroll_viewport`，两处入口共用一份。
        """
        canvas = getattr(self._scroll, "_parent_canvas", None)
        item = getattr(self._scroll, "_create_window_id", None)
        if canvas is None or item is None:  # 结构不符（版本差异）→ 退化为原行为
            return
        self._natural_h = self._scroll.winfo_height()  # 内容自然高，作为回落基准
        self._stretched_h = 0

        if self._natural_h > 0:
            canvas.bind("<Configure>", lambda _e: self._fit_scroll_viewport(), add="+")
            self.after_idle(self._fit_scroll_viewport)

    def _fit_scroll_viewport(self) -> None:
        """把内框高度对齐视口（唯一实现，``<Configure>`` 与 ``on_shown`` 都走这里）。"""
        # 尚未映射（启动早期）就不要动内框：此时 viewport 量不到真值，
        # 万一写进去会把滚动区弄坏。等 on_shown() / 下一个 <Configure> 再补。
        if not self.winfo_ismapped():
            return
        canvas = getattr(self._scroll, "_parent_canvas", None)
        item = getattr(self._scroll, "_create_window_id", None)
        if canvas is None or item is None:
            return
        try:
            viewport = canvas.winfo_height()
        except Exception:  # noqa: BLE001 - 拿不到尺寸就这一跳不处理
            return
        # 内容高于视口 → 不抬升，保持可滚动；视口尚未测量（==1）也跳过。
        if viewport <= 1 or viewport < self._natural_h:
            return
        if viewport == self._stretched_h:
            return
        self._stretched_h = viewport
        try:
            fraction = canvas.yview()[0]
        except Exception:  # noqa: BLE001
            fraction = 0.0
        canvas.itemconfigure(item, height=viewport)
        try:
            canvas.yview_moveto(fraction)
        except Exception:  # noqa: BLE001
            pass

    def on_shown(self) -> None:
        """页面被 tkraise 到最上层时由 ``app.show_page`` 调用（可选钩子）。

        常驻映射的副作用之一：切回本页时**不一定**伴随 ``<Configure>``
        （几何没变就不发），所以这里主动补算一次余高分配，否则窗口在别的页面上
        被改高/改矮之后再切回来，日志区会停在旧高度。
        ``_stretched_h`` 清零是关键 —— 它不是"已拉伸高度"而是"上次写进 item 的值"，
        不清零的话视口没变时 ``_fit_scroll_viewport`` 会直接 return，补算落空。
        """
        self._stretched_h = 0
        self.after_idle(self._fit_scroll_viewport)

    # ------------------------------------------------------------------ #
    # 2. 拖拽区
    # ------------------------------------------------------------------ #
    def _build_drop_zone(self) -> None:
        holder = ctk.CTkFrame(self._left, fg_color="transparent")
        holder.grid(row=1, column=0, sticky="nsew")
        holder.grid_columnconfigure(0, weight=1)
        holder.grid_rowconfigure(0, weight=1)

        self.drop_frame = ctk.CTkFrame(holder, corner_radius=th.RADIUS_XL,
                                       fg_color=th.DROP_FILL_IDLE)
        self.drop_frame.grid(row=0, column=0, sticky="nsew")
        self.drop_frame.grid_columnconfigure(0, weight=1)
        self.drop_frame.grid_rowconfigure(0, weight=1)

        canvas = self.drop_frame._canvas  # customtkinter 的内部画布
        canvas.configure(highlightthickness=0)
        canvas.bind("<Configure>", self._redraw_border)
        self._drop_state = "idle"        # 每次构建都从未选中开始
        self._drop_dim = True            # 尚未选择文件 → 不允许停在 selected
        self._drop_hovering = False

        inner = ctk.CTkFrame(self.drop_frame, fg_color="transparent")
        inner.place(relx=0.5, rely=0.5, anchor="center")
        self.drop_inner = inner

        self.drop_icon = ctk.CTkLabel(
            inner, text="⬇", font=th.font(24, "bold"), text_color=th.ACCENT
        )
        self.drop_icon.grid(row=0, column=0, pady=(0, 2))
        self.drop_title = ctk.CTkLabel(
            inner, text=DROP_TITLE_IDLE, font=th.font(14, "bold"),
            text_color=th.TEXT,
        )
        self.drop_title.grid(row=1, column=0)
        self.drop_hint = th.hint(inner, DROP_HINT_IDLE, anchor="center")
        self.drop_hint.grid(row=2, column=0, pady=(2, th.PAD_S))
        th.ghost_button(inner, "选择文件", self.choose_file, width=th.BTN_W_XL).grid(row=3, column=0)

        # 拖放接收：Windows 原生 WM_DROPFILES（详见 desktop/dnd.py 顶部说明）。
        # tkinterdnd2 的接收侧在本机实测完全收不到事件，故改为 ctypes 方案。
        # 回调由 dnd 内部用 Tk 定时器在主线程投递，这里直接操作控件即可。
        self.dnd_ok = dnd.hook(self, self._on_drop_files)
        if self.dnd_ok:
            dnd.set_hover_callback(
                lambda hovering: self.after_idle(
                    lambda: self._on_drag_hover(hovering)))
            self.app.log("拖放已启用：Windows 原生 WM_DROPFILES（可拖 .docx 到窗口任意位置）")
        else:
            # 退路：万一 ctypes 方案不可用，再试 tkinterdnd2（本机已知无效，仅兜底）
            self.dnd_ok = th.bind_drop_target(self.drop_frame, self._on_drop)
            self.app.log(f"WM_DROPFILES 不可用，改用 tkinterdnd2 注册 {th.last_registered_count()} 个控件")
            if not self.dnd_ok:
                self.drop_hint.configure(text="拖放组件不可用，请点下面的按钮选择文件")
        self.drop_frame.bind("<Button-1>", lambda _e: self.choose_file())

    def attach_motion(self, motion) -> None:
        """页面建好后由 app 注入轻量动效控制器，挂载拖拽区代码雨。"""
        self.rain = motion.attach_rain(
            self.drop_frame, self.drop_inner,
            on_click=self.choose_file, on_drop=self._on_drop,
        )
        self.rain.set_drop_idle(getattr(self, "_drop_state", "idle") == "idle")

    def _on_drag_hover(self, hovering: bool) -> None:
        """来自 dnd 的悬停通知：开始/结束一次"鼠标按住并停在窗口上"的过程。

        因为拿不到精确的 OLE 拖放事件，这里在按住期间自己持续判断光标是否
        落在拖拽区上方（``_track_drop_hover``），松手即结束。
        """
        self._drag_active = hovering
        if hovering:
            self._track_drop_hover()
        else:
            self._set_drop_hover(False)

    def _track_drop_hover(self) -> None:
        """按住期间每 80ms 复查一次光标位置，实时更新拖拽区外观。"""
        if not getattr(self, "_drag_active", False):
            return
        self._set_drop_hover(self._pointer_over_drop_zone())
        self.after(80, self._track_drop_hover)

    def _pointer_over_drop_zone(self) -> bool:
        """光标是否落在拖拽区上方（用于只在拖拽区高亮，而不是整个窗口）。"""
        try:
            x, y = self.drop_frame.winfo_pointerxy()
            widget = self.drop_frame.winfo_containing(x, y)
        except Exception:  # noqa: BLE001 - 取不到就当不在
            return False
        while widget is not None:
            if widget is self.drop_frame:
                return True
            widget = getattr(widget, "master", None)
        return False

    def _flash_step(self) -> None:
        """收到文件后让拖拽区"呼吸"两下，给一个明确的"收到了"反馈。

        常驻定时器（每 220ms 一跳）：``_flash_steps`` 由 ``_take_paths`` 置位，
        这里只在它 > 0 时真正呼吸；跳完最后一跳一定回落到 selected / idle
        （不能停在亮帧）。呼吸只切描边粗细（1px ↔ 3px），不切颜色——
        橙色在本 UI 里是"警告"，用在接收反馈上会让用户误以为出错。
        """
        self.after(220, self._flash_step)  # 先排下一次，这一跳无论如何都要续上
        if getattr(self, "_drop_state", "idle") != "pulse":
            return
        self._drop_target = self._current_spec()   # 跟随主题：每跳重取当前色对
        self._flash_steps -= 1
        self._flash_on = not self._flash_on
        self._drop_mix = 1.0                        # 脉冲是离散帧，不走插值
        self._drop_step = _DROP_TWEEN_STEPS
        self._drop_from = self._drop_target
        spec = _DROP_PULSE_SPECS["pulse_on" if self._flash_on else "pulse_off"]
        self._paint_frame(spec)
        if self._flash_steps <= 0:                  # 收尾：回落到常态
            self._flash_on = False
            self._apply_drop_state()

    def _current_spec(self) -> tuple:
        """当前逻辑态对应的参数表（pulse 取脉冲亮/暗帧）。"""
        state = getattr(self, "_drop_state", "idle")
        if state == "pulse":
            key = "pulse_on" if getattr(self, "_flash_on", False) else "pulse_off"
            return _DROP_PULSE_SPECS[key]
        return _DROP_STATES.get(state, _DROP_STATES["idle"])

    def _set_drop_hover(self, hovering: bool) -> None:
        """文件正被拖到窗口上时的视觉反馈（加粗变蓝 + 文案切换）。

        优先级高于 selected：拖新文件进来时压过"已选"外观。
        """
        hovering = bool(hovering)
        if hovering == getattr(self, "_drop_hovering", False) \
                and getattr(self, "_drop_state", "idle") in ("hover", "pulse"):
            return   # _track_drop_hover 每 80ms 会喂同一个值，别重复刷
        self._drop_hovering = hovering
        self._apply_drop_state()   # 悬停期间脉冲由 _apply_drop_state 作废（见那里的优先级注释）

    def _apply_drop_state(self) -> None:
        """按当前事实选定逻辑态并启动 120ms 过渡（状态表的唯一入口）。

        优先级：pulse 期间被悬停打断时**悬停优先**（拖新文件进来必须立刻看到
        "松手即可导入"），同时把没跑完的脉冲计数清零，避免它稍后再把态抢回去。
        """
        dim = getattr(self, "selected_file", None) is None
        self._drop_dim = dim
        was = getattr(self, "_drop_state", None)
        if was == "pulse" and getattr(self, "_drop_hovering", False):
            self._flash_steps = 0     # 脉冲被悬停打断：作废，不留在队列里等
            self._flash_on = False
        if getattr(self, "_flash_steps", 0) > 0:
            state = "pulse"
        elif getattr(self, "_drop_hovering", False):
            state = "hover"
        elif not dim:
            state = "selected"
        else:
            state = "idle"
        if state == was:
            if state == "selected":
                self._sync_drop_text()   # 文件没变，但换文件时文件名要跟上
            return
        self._drop_state = state
        rain = getattr(self, "rain", None)
        if rain is not None:
            rain.set_drop_idle(state == "idle")
        target = self._current_spec()
        self._start_drop_anim(target)

    def _start_drop_anim(self, target: tuple) -> None:
        """从当前外观插值到 ``target``：4 步 × 30ms（只用 after）。"""
        old = getattr(self, "_drop_target", target)
        # 起点 = 上一个态真正画出来的色对。中途被打断时，上一次插值可能只走了一半，
        # 这里不追那半步（视觉上无感），但 _paint_frame 立刻把 mix=0 画成 old，不会闪跳。
        # 约定：_drop_from 只放两个色对 (边框色对, 底色对)，与状态表的 7 元组索引不同。
        self._drop_from = (old[0], old[3])
        self._drop_target = target
        self._drop_step = 0
        self._drop_seq += 1
        seq = self._drop_seq
        self._drop_mix = 0.0
        self._paint_frame(target)
        self._animate_drop(seq)

    def _animate_drop(self, seq: int) -> None:
        """过渡帧回调；世代号不符（被新状态打断）就直接退出。"""
        if seq != getattr(self, "_drop_seq", -1):
            return
        try:
            if not self.drop_frame.winfo_exists():
                return
        except Exception:  # noqa: BLE001 - 控件已销毁
            return
        self._drop_step += 1
        self._drop_mix = min(1.0, self._drop_step / _DROP_TWEEN_STEPS)
        self._paint_frame(self._current_spec())
        if self._drop_step < _DROP_TWEEN_STEPS:
            self.after(_DROP_TWEEN_MS, lambda: self._animate_drop(seq))
        else:
            self._drop_seq += 1   # 让还在排队的旧帧彻底失效

    # --- 渲染 --------------------------------------------------------- #

    def _paint_frame(self, spec: tuple) -> None:
        """按 ``spec``（状态矩阵某一行）刷新拖拽区外观。

        ``_drop_mix`` < 1 时，边框色与底色都在 ``_drop_from`` → ``spec`` 之间插值。
        **两套索引不要混**：``spec`` 是状态表的 7 元组（0 边框/3 底/1 线宽/2 gap…），
        而 ``_drop_from`` 只存两个色对 ``(边框色对, 底色对)``，所以这里读 ``src[1]``。
        ``CTkFrame`` 的底色吃色对（切主题跟随），底层 ``_canvas`` 的 bg 只吃单色，
        所以这里用 theme.C() 解一次——与 theme.C 的既有用法一致。
        """
        frame = getattr(self, "drop_frame", None)
        if frame is None:
            return
        mix = getattr(self, "_drop_mix", 1.0)
        src = getattr(self, "_drop_from", spec)
        if mix < 1 and len(src) >= 2:
            border = _blend_spec(src[0], spec[0], mix)
            fill = _blend_spec(src[1], spec[3], mix)
        else:
            border, fill = spec[0], spec[3]
        self._drop_bg_frame = fill
        self._drop_bg_solid = th.C(border)
        self._paint_drop_bg(frame, fill)
        self._sync_drop_text(spec)
        self._render_drop_border(th.C(border), spec[1], spec[2])

    def _paint_drop_bg(self, frame, fill) -> None:
        """写底色：CTkFrame 吃色对，_canvas 的 bg 只能吃单色（兜底用上一次的）。"""
        try:
            frame.configure(fg_color=fill)
        except Exception:  # noqa: BLE001 - 控件正在销毁
            return
        try:
            current = frame.cget("fg_color")
        except Exception:  # noqa: BLE001
            return
        color = current[0] if isinstance(current, (tuple, list)) else current
        if not isinstance(color, str) or not _hex_to_rgb(color):
            color = getattr(self, "_drop_bg_solid", color)   # 'transparent' / 'gray95' 解析不了
        try:
            frame._canvas.configure(bg=color)
        except Exception:  # noqa: BLE001
            pass
        self._drop_bg_solid = color

    def _render_drop_border(self, color: str, line_width: float, gap: int) -> None:
        """画边框：``gap <= 1`` 是实线（selected 态），否则走 theme 的虚线画法。"""
        frame = getattr(self, "drop_frame", None)
        if frame is None:
            return
        canvas = frame._canvas
        width, height = frame.winfo_width(), frame.winfo_height()
        if width < 2 or height < 2:   # 布局初期会给出 1px，画了就是退化虚线
            return
        self._drop_line = color
        self._drop_dash = (gap, line_width)
        if gap <= 1:
            canvas.delete("dashborder")
            canvas.create_rectangle(
                1, 1, max(width - 2, 2), max(height - 2, 2),
                outline=color, width=line_width, fill="",
                tags=("dashborder", "_dnd_border"))
        else:
            th.dashed_border(canvas, width, height, color=color, gap=gap, line_width=line_width)

    def _sync_drop_text(self, spec: Optional[tuple] = None) -> None:
        """文字/图标立刻切到目标态（只插值颜色，不插值文字）。"""
        if spec is None:
            spec = self._current_spec()
        icon, title, hint = spec[4], spec[5], spec[6]
        if hint is None:   # selected：带上已选文件名
            name = getattr(getattr(self, "selected_file", None), "name", "")
            if len(name) > 28:
                name = name[:26] + "…"
            hint = DROP_HINT_SELECTED_FMT.format(name or "—")
        title_label = getattr(self, "drop_title", None)
        if title_label is None:
            return
        title_label.configure(text=title)
        self.drop_hint.configure(text=hint)
        self.drop_icon.configure(text=icon)

    def _redraw_border(self, event=None) -> None:
        """``<Configure>`` 与外部重绘入口：按当前态 + 当前插值位置重画一遍。"""
        if getattr(self, "_drop_seq", None) is None:
            return
        self._drop_seq += 1   # 尺寸变了，把在飞的动画帧作废，避免它按旧尺寸重画
        self._paint_frame(self._current_spec())

    def _on_drop(self, event) -> None:
        """tkinterdnd2 的 <<Drop>> 事件（兜底路径，本机已知收不到）。"""
        paths = th.parse_drop_event(event)
        self._take_paths(paths)

    def _on_drop_files(self, paths: List[str]) -> None:
        """WM_DROPFILES 回调：由 dnd.py 的 Tk 定时器在主线程调用。"""
        print(f"[dnd] 收到拖放：{paths}", file=sys.stderr)
        self._take_paths(list(paths))

    def _take_paths(self, paths: List[str]) -> None:
        self._set_drop_hover(False)  # 拖放结束，恢复常态外观
        if not paths:
            self.app.status("拖入内容无法识别为文件路径")
            return
        self._flash_steps = 4  # 呼吸两下：确认"收到了"
        self._apply_drop_state()   # 立刻切到 pulse 态（不改动 load_file 的业务逻辑）
        self.app.log(f"拖入 {len(paths)} 个文件：{'、'.join(Path(p).name for p in paths)}")
        self.load_file(paths[0], from_drop=True)
        if len(paths) > 1:
            self.app.status(f"已接收 {Path(paths[0]).name}，忽略另外 {len(paths) - 1} 个文件")

    # ------------------------------------------------------------------ #
    # 3. 两列内容：文件卡进左列，参数卡进右列
    # ------------------------------------------------------------------ #

    def _build_columns(self) -> None:
        """左列 = 已选文件卡；右列 = 处理参数卡。（两列容器在 __init__ 里已建好）"""
        self._build_file_card(self._left)
        self._build_option_card(self._right)

    def _build_file_card(self, master) -> None:
        card = th.card(master)
        card.grid(row=2, column=0, sticky="new")
        card.grid_columnconfigure(0, weight=1)

        th.section_title(card, "已选文件").grid(
            row=0, column=0, sticky="w", padx=th.PAD, pady=(th.PAD, 2)
        )
        self.file_name = ctk.CTkLabel(
            card, text="尚未选择文件", font=th.font(13), anchor="w",
            wraplength=420, justify="left",
        )
        self.file_name.grid(row=1, column=0, sticky="ew", padx=th.PAD, pady=(0, th.PAD_S))

        meta = ctk.CTkFrame(card, fg_color="transparent")
        meta.grid(row=2, column=0, sticky="ew", padx=th.PAD, pady=(0, th.PAD))
        meta.grid_columnconfigure(1, weight=1)
        self.file_meta_labels = {}
        for i, key in enumerate(("大小", "段落数", "修改时间")):
            ctk.CTkLabel(meta, text=key, font=th.font(12), anchor="w",
                         text_color=th.TEXT_SECOND).grid(
                row=i, column=0, sticky="w", pady=1)
            value = ctk.CTkLabel(meta, text="—", font=th.font(12), anchor="w")
            value.grid(row=i, column=1, sticky="w", padx=(th.PAD_S, 0), pady=1)
            self.file_meta_labels[key] = value

    def _build_option_card(self, master) -> None:
        card = th.card(master)
        card.grid(row=0, column=0, sticky="new")
        card.grid_columnconfigure(0, weight=1)

        th.section_title(card, "处理参数").grid(
            row=0, column=0, sticky="w", padx=th.PAD, pady=(th.PAD, 2)
        )

        ctk.CTkLabel(card, text="格式要求预设", font=th.font(12), anchor="w").grid(
            row=1, column=0, sticky="w", padx=th.PAD, pady=(th.PAD_S, 2)
        )
        self.preset_var = tk.StringVar(value="")
        self.preset_menu = th.combo(
            card, values=["（加载中…）"], variable=self.preset_var,
            command=self._on_preset_change, height=34,
        )
        self.preset_menu.grid(row=2, column=0, sticky="ew", padx=th.PAD)

        th.ghost_button(card, "上传格式要求…", self.upload_format, height=30).grid(
            row=3, column=0, sticky="ew", padx=th.PAD, pady=(th.PAD_S, th.PAD_S)
        )

        ctk.CTkLabel(card, text="输出目录", font=th.font(12), anchor="w").grid(
            row=4, column=0, sticky="w", padx=th.PAD, pady=(th.PAD_S, 2)
        )
        row = ctk.CTkFrame(card, fg_color="transparent")
        row.grid(row=5, column=0, sticky="ew", padx=th.PAD)
        row.grid_columnconfigure(0, weight=1)
        self.out_var = tk.StringVar(value=self.app.settings.output_dir)
        th.entry(row, textvariable=self.out_var).grid(row=0, column=0, sticky="ew")
        th.ghost_button(row, "浏览", self.choose_output_dir, width=56, height=34).grid(
            row=0, column=1, padx=(th.PAD_S, 0)
        )
        th.hint(card, "默认存到「文档\\Taofu」，改完立即记住（设置页可看完整路径）").grid(
            row=6, column=0, sticky="w", padx=th.PAD, pady=(4, th.PAD)
        )
        self.out_var.trace_add("write", lambda *_: self.app.set_output_dir(self.out_var.get()))

    # ------------------------------------------------------------------ #
    # 4. 按钮行
    # ------------------------------------------------------------------ #

    def _build_actions(self) -> None:
        bar = ctk.CTkFrame(self._right, fg_color="transparent")
        bar.grid(row=1, column=0, sticky="ew", pady=(th.PAD, 0))
        bar.grid_columnconfigure(0, weight=1)

        right = ctk.CTkFrame(bar, fg_color="transparent")
        right.grid(row=0, column=1, sticky="e")
        self.preview_button = th.ghost_button(
            right, "预览效果", self.preview, width=th.BTN_W_XL, height=38
        )
        self.preview_button.grid(row=0, column=0, padx=(0, th.PAD_S))
        process_wrap = ctk.CTkFrame(right, fg_color="transparent")
        process_wrap.grid(row=0, column=1)
        process_wrap.grid_columnconfigure(1, weight=1)
        edge = ctk.CTkFrame(
            process_wrap, width=2, height=th.CTRL_H_L,
            fg_color=th.ACCENT, corner_radius=0,
        )
        edge.grid(row=0, column=0, sticky="ns")
        self.process_button = th.primary_button(
            process_wrap, "< START_PROCESS >", self.start_process,
            width=170, font=th.mono_font(11, "bold"),
        )
        self.process_button.grid(row=0, column=1, padx=(2, 0))
        self._process_hovered = False

        def hover_in(_event=None) -> None:
            self._process_hovered = True
            if not getattr(self, "_busy", False):
                self.process_button.configure(text="< //START_PROCESS// >")

        def hover_out(_event=None) -> None:
            self._process_hovered = False
            if not getattr(self, "_busy", False):
                self.process_button.configure(text="< START_PROCESS >")

        for target in (self.process_button,
                       getattr(self.process_button, "_canvas", None)):
            if target is not None:
                target.bind("<Enter>", hover_in, add="+")
                target.bind("<Leave>", hover_out, add="+")

    # ------------------------------------------------------------------ #
    # 5. 日志区
    # ------------------------------------------------------------------ #

    def _build_log(self) -> None:
        wrap = ctk.CTkFrame(self._left, fg_color="transparent")
        wrap.grid(row=3, column=0, sticky="nsew", pady=(th.PAD, th.PAD_L))
        wrap.grid_columnconfigure(0, weight=1)
        wrap.grid_rowconfigure(1, weight=1)

        head = ctk.CTkFrame(wrap, fg_color="transparent")
        head.grid(row=0, column=0, sticky="ew")
        head.grid_columnconfigure(0, weight=1)

        self.log_toggle = ctk.CTkButton(
            head, text=self._log_title(), font=th.mono_font(12, "bold"), anchor="w",
            fg_color="transparent", hover=False, text_color=th.TEXT,
            command=self.toggle_log,
        )
        self.log_toggle.grid(row=0, column=0, sticky="w")
        th.ghost_button(head, "清空", self.clear_log, width=56, height=26).grid(
            row=0, column=1, sticky="e"
        )

        # 日志框走工厂：等宽字（Consolas）+ LOG_FONT_SIZE 仍是原来的口径，
        # 底色/描边/圆角由 th.textbox 收口。
        self.log_box = th.textbox(wrap, mono=True, height=150)
        self.log_box.grid(row=1, column=0, sticky="nsew", pady=(th.PAD_S, 0))
        self.log_box.configure(state="disabled")
        text = self.log_box._textbox
        text.tag_configure("stamp", foreground=th.C(th.TEXT_MUTED))
        text.tag_configure("INFO", foreground=th.LOG_INFO)
        text.tag_configure("SYS", foreground=th.C(th.ACCENT))
        text.tag_configure("WARN", foreground=th.LOG_WARN)
        text.tag_configure("ERR", foreground=th.C(th.SECONDARY))
        self.log_prompt = ctk.CTkLabel(
            wrap, text="root@taofu:~# █", font=th.mono_font(10),
            text_color=th.ACCENT_TEXT, anchor="w",
        )
        self.log_prompt.grid(row=2, column=0, sticky="w", pady=(4, 0))
        self._log_visible = bool(self.app.settings.log_expanded)
        self._apply_log_visibility()

    def _log_title(self) -> str:
        mark = "▾ " if getattr(self, "_log_visible", True) else "▸ "
        return mark + "[ SYSTEM_LOG ]"

    def _apply_log_visibility(self) -> None:
        if self._log_visible:
            self.log_box.grid()
            if hasattr(self, "log_prompt"):
                self.log_prompt.grid()
        else:
            self.log_box.grid_remove()
            if hasattr(self, "log_prompt"):
                self.log_prompt.grid_remove()
        self.log_toggle.configure(text=self._log_title())

    def toggle_log(self) -> None:
        self._log_visible = not getattr(self, "_log_visible", True)
        self._apply_log_visibility()
        self.app.settings.log_expanded = self._log_visible

    @staticmethod
    def _log_level(line: str) -> str:
        raw = line.strip()
        upper = raw.upper()
        if ("✕" in raw or "[ERR]" in upper or "失败" in raw
                or "错误" in raw or "ERROR" in upper or "TRACEBACK" in upper):
            return "ERR"
        if "⚠" in raw or "[WARN]" in upper or "警告" in raw:
            return "WARN"
        if (raw.startswith("=") or raw.startswith("桃符 v")
                or "开始处理" in raw or raw.startswith("预设：")
                or raw.startswith("预设标识：") or raw.startswith("输出：")):
            return "SYS"
        return "INFO"

    def append_log(self, text: str) -> None:
        lines = text.splitlines() or [""]
        self.log_box.configure(state="normal")
        for line in lines:
            stamp = time.strftime("%H:%M:%S")
            level = self._log_level(line)
            self.log_box.insert("end", f"[{stamp}] ", "stamp")
            self.log_box.insert("end", f"[{level}] ", level)
            self.log_box.insert("end", line + "\n")
        self.log_box.see("end")
        self.log_box.configure(state="disabled")

    def clear_log(self) -> None:
        self.log_box.configure(state="normal")
        self.log_box.delete("1.0", "end")
        self.log_box.configure(state="disabled")

    # ------------------------------------------------------------------ #
    # 业务动作（Phase 3：真调插件 CLI）
    # ------------------------------------------------------------------ #

    def refresh_presets(self, presets=None) -> None:
        """把「内置 + 用户」预设灌进下拉框，展示名带 📦/👤 前缀。"""
        items: List[pl.Preset] = list(self.app.presets if presets is None else presets)
        self._display_to_name = {p.display_name: p.name for p in items}
        displays = [p.display_name for p in items]
        if not displays:
            self.preset_menu.configure(values=["（没有可用预设）"])
            self.preset_var.set("")
            self._refresh_steps()
            return
        self.preset_menu.configure(values=displays)

        current = self._display_to_name.get(self.preset_var.get(), "")
        if current not in self._display_to_name.values():
            preferred = self.app.settings.last_preset
            names = {p.name for p in items}
            chosen = preferred if preferred in names else items[0].name
            current = chosen
        for display, name in self._display_to_name.items():
            if name == current:
                self.preset_var.set(display)
                break
        self.app.set_last_preset(current)
        self._refresh_steps()   # 启动时走的就是这条路径：第 2 节点的初值在这里定

    def current_preset(self) -> Optional[pl.Preset]:
        """当前选中的预设对象（找不到返回 None）。"""
        raw = self.preset_var.get()
        name = self._display_to_name.get(raw, raw)
        return pl.find_preset(name, self.app.presets)

    def current_preset_name(self) -> str:
        raw = self.preset_var.get()
        return self._display_to_name.get(raw, raw)

    def current_presets_dir(self) -> Optional[str]:
        """当前预设该带的 ``--presets-dir``（内置 → None，用户 → 用户目录）。"""
        return pl.presets_dir_arg(self.current_preset())

    def current_cli_name(self) -> str:
        """插件 CLI 认的预设标识（用户预设 = 文件名去后缀）。"""
        preset = self.current_preset()
        if preset is not None:
            return preset.cli_name
        return self.current_preset_name()

    def _on_preset_change(self, choice: str) -> None:
        name = self._display_to_name.get(choice, choice)
        preset = pl.find_preset(name, self.app.presets)
        source = preset.source_label if preset else "未知来源"
        location = "用户预设目录" if preset is not None and preset.is_user else "presets/"
        self.app.set_last_preset(name)
        self.app.status(f"已选择预设：{name}（{source}）")
        self.app.log(f"预设切换 → {name}（{source} · {location}）")
        self._processed_ok = False   # 换了预设，上一轮的结果不再对应当前配置
        self._refresh_steps()

    def choose_file(self) -> None:
        path = filedialog.askopenfilename(
            title="选择要处理的 Word 文档",
            filetypes=[("Word 文档", "*.docx"), ("所有文件", "*.*")],
            parent=self.app,
        )
        if path:
            self.load_file(path)

    def load_file(self, path: str | Path, from_drop: bool = False) -> None:
        target = Path(path)
        self.selected_file = target
        self._processed_ok = False   # 换了文件，上一轮的结果不再对应当前输入
        self._refresh_steps()
        self.file_name.configure(text=target.name)
        if from_drop:
            self.app.status(f"已接收拖入的文件：{target.name} · 正在读取信息…")

        size_label = "—"
        try:
            stat = target.stat()
            size_label = cfg.human_size(stat.st_size)
            mtime = cfg.fmt_mtime(target)
        except OSError as exc:
            self.file_meta_labels["大小"].configure(text="读取失败")
            self.file_meta_labels["段落数"].configure(text="—")
            self.file_meta_labels["修改时间"].configure(text="—")
            self.app.log(f"文件读取失败：{exc}")
            self.app.status(f"拖入的文件读不出来：{target.name}（{exc}）")
            return

        self.file_meta_labels["大小"].configure(text=size_label)
        self.file_meta_labels["修改时间"].configure(text=mtime)
        self.file_meta_labels["段落数"].configure(text="读取中…")

        self.app.status(f"已选择文件：{target.name}")
        self.app.log(f"已选择文件 {target}（{size_label}）")

        if target.suffix.lower() != DOCX_SUFFIX:
            self.file_meta_labels["段落数"].configure(text="非 .docx")
            self.app.log("注意：扩展名不是 .docx，处理只支持 Word 文档")
            self.app.status(f"已接收 {target.name}，但它不是 .docx，无法处理")
        else:
            self.after(30, self._count_paragraphs)

    def _count_paragraphs(self) -> None:
        """用 python-docx 数段落（只读）。失败不影响其它信息。"""
        target = self.selected_file
        if target is None:
            return
        try:
            from docx import Document

            document = Document(str(target))
            total = len(document.paragraphs)
            nonempty = sum(1 for p in document.paragraphs if p.text.strip())
            self.file_meta_labels["段落数"].configure(text=f"{total}（非空 {nonempty}）")
            self.app.log(f"段落统计：{total} 段，非空 {nonempty} 段，表格 {len(document.tables)} 个")
            if total:
                self.app.status(f"已就绪：{target.name} · {total} 段 · 可点「开始处理」")
        except Exception as exc:  # noqa: BLE001 - 保护性：坏文件不该弄崩界面
            self.file_meta_labels["段落数"].configure(text="读取失败")
            self.app.log(f"段落统计失败：{type(exc).__name__}: {exc}")

    def choose_output_dir(self) -> None:
        path = filedialog.askdirectory(
            title="选择输出目录", initialdir=self.out_var.get() or None, parent=self.app
        )
        if path:
            self.out_var.set(path)

    def upload_format(self) -> None:
        """上传格式要求：解析成 DSL，可保存为用户预设。"""
        result = up_dlg.open_upload_dialog(self.app)
        if result is None:
            self.app.log("上传格式要求：已取消")
            return
        self.app.log(f"上传格式要求：{result.summary}")
        if result.saved_path is not None:
            self.app.reload_presets(select=result.name)
            self.app.status(f"已新增用户预设「{result.name}」")
        elif result.name:
            self.app.status(f"本次使用上传的格式要求：{result.name}")

    def preview(self) -> None:
        """方案二：真跑一次处理到临时目录，按真实输出做预览。"""
        if not self._require_ready():
            return
        preset = self.current_preset()
        self.app.log(
            f"预览：用预设「{self.current_preset_name()}」试处理 {self.selected_file.name}…"
        )
        self.app.status("正在生成预览（临时目录中真实处理一次）…")
        self.app.progress_mode("indeterminate")
        try:
            self.update_idletasks()
            chosen = pv.show_preview(
                self.app,
                self.selected_file,
                self.current_preset_name(),
                self.out_var.get(),
                presets_dir=self.current_presets_dir(),
                on_confirm=self.start_process,
                log=self.app.log,
            )
        finally:
            self.app.progress_mode("determinate")
        if chosen:
            self.app.status("预览完成 · 可直接点「开始处理」")
        else:
            self.app.status("预览已关闭")

    def start_process(self) -> None:
        """真调 bridge.process：后台线程干活，主线程轮询取结果。

        刻意**不在子线程里碰 Tk**：``after`` 的跨线程调度在 Windows 上不保证送达
        （实测回调会丢）。子线程只写 ``self._worker_result`` / ``self._worker_error``，
        主线程的 :meth:`_poll_process` 取走并回填界面。
        """
        if self._busy:
            self.app.status("正在处理中，请稍候…")
            return
        # 点击主按钮：一次 100ms 的颜色 glitch，不阻塞后续校验/处理。
        self.app.motion.flash(self.process_button, "border_color")
        if not self._require_ready():
            return

        preset_name = self.current_cli_name()
        preset_label = self.current_preset_name()
        presets_dir = self.current_presets_dir()
        target = self.selected_file
        out_dir = (self.out_var.get() or "").strip() or self.app.settings.output_dir
        if not cfg.ensure_dir(out_dir):
            dlg.error(self.app, f"输出目录不可用：{out_dir}", title="无法处理")
            return
        output = self._unique_output_path(Path(out_dir), target)

        self._worker_result = None
        self._worker_error = None
        self._processed_ok = False   # 新一轮开始：第 3 节点熄灭
        self._refresh_steps()
        self._set_busy(True)
        self.app.status(f"正在处理 {target.name} …")
        self.app.progress_mode("indeterminate")
        self.app.log("=" * 56)
        self.app.log(f"开始处理：{target.name}")
        self.app.log(f"预设：{preset_label}（{'用户预设目录' if presets_dir else 'presets/'}）")
        self.app.log(f"预设标识：{preset_name}")
        self.app.log(f"输出：{output}")

        def worker() -> None:
            try:
                self._worker_result = bridge.process(
                    str(target), preset_name, str(output), presets_dir=presets_dir
                )
            except bridge.BridgeError as exc:
                self._worker_error = (str(exc), exc.detail)
            except Exception as exc:  # noqa: BLE001 - 兜底，别把线程炸了
                self._worker_error = (f"{type(exc).__name__}: {exc}", "")

        threading.Thread(target=worker, name="wfa-process", daemon=True).start()
        self.after(120, self._poll_process)

    def _poll_process(self) -> None:
        """主线程轮询：子线程写结果，这里读走并回填界面。"""
        error = getattr(self, "_worker_error", None)
        if error is not None:
            self._worker_error = None
            self._on_process_failed(error[0], error[1])
            return
        summary = getattr(self, "_worker_result", None)
        if summary is not None:
            self._worker_result = None
            self._on_process_done(summary)
            return
        self.after(120, self._poll_process)

    # -- 处理结果回调（都在主线程） ------------------------------------- #

    def _relocate_report(self, report: str, out_docx: str) -> str:
        """把临时目录里的报告移到输出文件旁边；移不动就保留临时路径。"""
        if not report or not out_docx:
            return report
        source = Path(report)
        if not source.is_file():
            return report          # 报告本来就没写出来，别再报一次移动失败
        try:
            source.resolve().relative_to(Path(tempfile.gettempdir()).resolve())
        except (OSError, ValueError):
            return report          # 不是临时目录里的（老版本/自定义路径），保持原样
        destination = Path(out_docx + ".report.json")
        try:
            destination.parent.mkdir(parents=True, exist_ok=True)
            if destination.exists():
                destination.unlink()
            shutil.move(str(source), str(destination))
            self.app.log(f"报告已移到输出目录：{destination}")
            return str(destination)
        except OSError as exc:
            self.app.log(f"[WARN] 报告文件移动失败，已跳过：{type(exc).__name__}: {exc}")
            return report

    def _on_process_done(self, summary: dict) -> None:
        self._set_busy(False)
        self.app.progress_mode("determinate")
        for line in summary.get("_stderr", "").splitlines():
            if line.strip():
                self.app.log(line.rstrip())
        ok = bool(summary.get("ok"))
        self._processed_ok = ok
        self._refresh_steps()
        applied = summary.get("applied", 0)
        skipped = summary.get("skipped", 0)
        out_docx = summary.get("output_docx") or ""
        report = ""
        if ok:
            # 只有真成功才去“收养”报告；失败时报告可能压根没生成，
            # 硬移一次只会在日志里再刷一条 PermissionError 干扰判断。
            report = self._relocate_report(summary.get("report_json") or "", out_docx)
        else:
            report = summary.get("report_json") or ""

        if ok:
            self.app.log(f"完成：改动 {applied} 段，跳过 {skipped} 段，"
                         f"文本{'未变' if summary.get('text_unchanged') else '已变'}")
            cleanup = summary.get("cleanup") or {}
            if cleanup:
                self.app.log(f"文本清理：{cleanup.get('chars_removed', 0)} 字符 / "
                             f"{cleanup.get('paragraphs_removed', 0)} 段被删")
            structure = summary.get("structure") or {}
            if structure:
                self.app.log(f"结构预处理：识别标题 {structure.get('headings_mapped', 0)} 个 "
                             f"（信号 {'、'.join(structure.get('signals') or []) or '无'}）")
            for warn in summary.get("warnings") or []:
                self.app.log(f"⚠ {warn}")
            self.app.log(f"报告：{report}")
            self.app.status(f"处理完成 · {Path(out_docx).name if out_docx else '已输出'}")
            detail = (f"输出：{out_docx}\n报告：{report}\n"
                      f"改动段落：{applied} · 跳过：{skipped}\n"
                      f"LLM 调用：{summary.get('llm_calls', 0)} 次")
            dlg.info(self.app, "处理完成。", title="处理成功", detail=detail)
        else:
            err = summary.get("_error") or "处理失败"
            headline = err.splitlines()[0].strip() if err.strip() else "处理失败"
            self.app.log(f"✕ 处理失败：{headline}")
            for line in err.splitlines()[1:]:
                if line.strip():
                    self.app.log(f"  {line.rstrip()}")
            self.app.status("处理失败")
            self.app.motion.flash(self.app.status_label, "text_color")
            dlg.error(self.app, headline, title="处理失败", detail=err)

    def _on_process_failed(self, message: str, detail: str) -> None:
        self._set_busy(False)
        self.app.progress_mode("determinate")
        self._processed_ok = False   # 失败不点亮第 3 节点
        self._refresh_steps()
        self.app.log(f"✕ {message}")
        self.app.status("处理失败")
        self.app.motion.flash(self.app.status_label, "text_color")
        dlg.error(self.app, message, title="处理失败", detail=detail)

    # -- 小工具 ---------------------------------------------------------- #

    def _set_busy(self, busy: bool) -> None:
        self._busy = busy
        try:
            if busy:
                text = "< PROCESSING >"
            elif getattr(self, "_process_hovered", False):
                text = "< //START_PROCESS// >"
            else:
                text = "< START_PROCESS >"
            self.process_button.configure(
                state="disabled" if busy else "normal",
                text=text,
            )
            self.preview_button.configure(state="disabled" if busy else "normal")
        except Exception:  # noqa: BLE001 - 控件已销毁时忽略
            pass

    def _require_ready(self) -> bool:
        if self.selected_file is None:
            dlg.info(self.app, "请先选择一个 .docx 文件。", title="还没有文件")
            return False
        if self.selected_file.suffix.lower() != DOCX_SUFFIX:
            dlg.error(self.app, "只支持 .docx 文件。", title="格式不支持",
                      detail=str(self.selected_file))
            return False
        if self.current_preset() is None:
            dlg.error(self.app, "请选择一个格式要求预设。", title="还没有预设")
            return False
        return True

    @staticmethod
    def _unique_output_path(out_dir: Path, target: Path) -> Path:
        """``<stem>_formatted.docx``；已存在时加 ``-2``、``-3`` 后缀。"""
        stem = target.stem or "output"
        candidate = out_dir / f"{stem}_formatted.docx"
        if not candidate.exists():
            return candidate
        for i in range(2, 1000):
            candidate = out_dir / f"{stem}_formatted-{i}.docx"
            if not candidate.exists():
                return candidate
        return candidate
