"""Tests for the opt-in text cleanup pass (markdown residue + junk symbols).

The cleanup pass is the only code path allowed to change document text, so it is
tested in isolation (``TextCleaner``), at the parser boundary (``清理：`` lines)
and end to end (a markdown-polluted .docx through ``apply_format``).
"""

from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "src")
for path in (SRC, ROOT):
    if path not in sys.path:
        sys.path.insert(0, path)

import docx  # noqa: E402
from docx.enum.text import WD_BREAK  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt  # noqa: E402

from word_format_applier import applier, cleanup, parser, textparse  # noqa: E402

ALL_ON = {
    "enabled": True, "markdown": True, "invisible": True, "punctuation": True,
    "spaces": True, "decorations": True, "blank_lines": True,
    "keep_symbols": "%‰°±",
}

ZWSP = "\u200b"
ZWNJ = "\u200c"
BOM = "\ufeff"


def cleaner(**overrides):
    config = dict(ALL_ON)
    config.update(overrides)
    return cleanup.TextCleaner(config)


class TestCleanerBasics(unittest.TestCase):

    def test_markdown_headings(self):
        clean = cleaner()
        self.assertEqual(clean.clean("# 一级标题")[0], "一级标题")
        self.assertEqual(clean.clean("## 二级标题 ##")[0], "二级标题")
        self.assertEqual(clean.clean("# 一级标题")[1]["heading"], 1)
        self.assertEqual(clean.clean("###### 六级")[1]["heading"], 6)
        # no space after the hashes is not markdown, but the hashes are still noise
        self.assertEqual(clean.clean("###无空格")[0], "无空格")

    def test_markdown_lists_and_quotes(self):
        clean = cleaner()
        self.assertEqual(clean.clean("- 列表项一")[0], "列表项一")
        self.assertEqual(clean.clean("* 星号列表")[0], "星号列表")
        self.assertEqual(clean.clean("1. 有序列表")[0], "有序列表")
        self.assertEqual(clean.clean("2) 括号编号")[0], "括号编号")
        self.assertEqual(clean.clean("- [x] 已完成任务")[0], "已完成任务")
        self.assertEqual(clean.clean("> 引用内容")[0], "引用内容")
        self.assertEqual(clean.clean(">> 嵌套引用")[0], "嵌套引用")

    def test_markdown_inline(self):
        clean = cleaner()
        self.assertEqual(clean.clean("**加粗文字** 与 *斜体*")[0], "加粗文字 与 斜体")
        self.assertEqual(clean.clean("__双下划线__")[0], "双下划线")
        self.assertEqual(clean.clean("这是 `代码` 片段")[0], "这是 代码 片段")
        self.assertEqual(clean.clean("[链接文字](http://example.com)")[0], "链接文字")
        self.assertEqual(clean.clean("![图片说明](img.png)")[0], "图片说明")
        self.assertEqual(clean.clean("~~删除线~~")[0], "删除线")
        # snake_case must survive the single-underscore emphasis rule
        self.assertEqual(clean.clean("变量 file_name_here 保留")[0], "变量 file_name_here 保留")

    def test_markdown_rules_and_fences(self):
        clean = cleaner()
        text, info = clean.clean("---")
        self.assertEqual(text, "")
        self.assertEqual(info["setext"], 2)
        self.assertTrue(info["blank"])
        text, info = clean.clean("===")
        self.assertEqual(info["setext"], 1)

        clean = cleaner()
        self.assertEqual(clean.clean("```python")[0], "")
        # fenced content is code: it must come through verbatim
        self.assertEqual(clean.clean("print('a *b* c')")[0], "print('a *b* c')")
        self.assertEqual(clean.clean("```")[0], "")

    def test_markdown_tables(self):
        clean = cleaner()
        self.assertEqual(clean.clean("| 指标 | 数值 |")[0], "指标 数值")
        self.assertEqual(clean.clean("| --- | --- |")[0], "")

    def test_invisible_characters(self):
        clean = cleaner()
        self.assertEqual(clean.clean("零%s宽%s字%s符" % (ZWSP, ZWNJ, BOM))[0], "零宽字符")

    def test_punctuation_and_spaces(self):
        clean = cleaner()
        self.assertEqual(clean.clean("复，，句。。号！！")[0], "复，句。号！")
        self.assertEqual(clean.clean("全角\u3000\u3000空格")[0], "全角 空格")
        self.assertEqual(clean.clean("  首尾空格  ")[0], "首尾空格")
        # … and —— are real Chinese punctuation, not duplication
        self.assertEqual(clean.clean("等等…… 破折——号")[0], "等等…… 破折——号")

    def test_decoration_symbols_are_contextual(self):
        clean = cleaner()
        self.assertEqual(clean.clean("■ 装饰项目")[0], "装饰项目")
        self.assertEqual(clean.clean("◆◆◆ 占位")[0], "占位")
        self.assertEqual(clean.clean("■■■")[0], "")
        self.assertEqual(clean.clean("· 项目符号")[0], "项目符号")
        # a single mid-sentence separator is meaningful and survives
        self.assertEqual(clean.clean("· 张三·李四")[0], "张三·李四")

    def test_meaningful_symbols_survive(self):
        clean = cleaner()
        for text in ("增长率 12% 达成", "温度 25°C", "误差 ±3", "比例 5‰",
                     "2026 年度工作总结报告", "一、总体情况", "1.5倍行距说明",
                     "-5℃ 低温", "3 + 4 = 7"):
            self.assertEqual(clean.clean(text)[0], text, text)

    def test_prose_punctuation_is_not_treated_as_orphan_junk(self):
        clean = cleaner()
        for text in ("本次会议要求如下：", "（一）总体要求：", "他说——",
                     "见 2020-2021 年度报告", "等等……", "备注："):
            self.assertEqual(clean.clean(text)[0], text, text)
        # but a dangling list separator really is junk
        self.assertEqual(clean.clean("这是一个悬空的、")[0], "这是一个悬空的")
        self.assertEqual(clean.clean("尾部残留--")[0], "尾部残留")


class TestCleanupDisabling(unittest.TestCase):

    def test_no_categories_means_no_cleanup(self):
        clean = cleanup.TextCleaner({})
        self.assertFalse(clean.enabled)
        self.assertEqual(clean.clean("# 标题 **粗**")[0], "# 标题 **粗**")

    def test_markdown_only_leaves_symbols_alone(self):
        clean = cleaner(decorations=False, punctuation=False, spaces=False)
        self.assertEqual(clean.clean("# 标题，，■")[0], "标题，，■")

    def test_keep_symbols_protects_decoration(self):
        clean = cleaner(keep_symbols="■")
        self.assertEqual(clean.clean("■ 装饰项目")[0], "■ 装饰项目")
        self.assertEqual(clean.clean("◆ 另一个")[0], "另一个")

    def test_remove_symbols_adds_targets(self):
        clean = cleaner(remove_symbols="§¶")
        self.assertEqual(clean.clean("§ 段落")[0], "段落")


class TestCleanupDirectives(unittest.TestCase):

    def test_directive_is_not_read_as_a_font_clause(self):
        rules = textparse.parse_text_rules(
            "清理：markdown 不可见字符 重复标点 空白 装饰符号 碎片 空段落 删除空段落\n"
            "保留符号：% ‰ ° ±\n"
            "标题：黑体 三号 居中 加粗\n"
            "正文：宋体 小四 1.5倍行距 首行缩进2字符 两端对齐\n")
        block = rules["cleanup"]
        self.assertTrue(block["enabled"])
        for key in cleanup.CLEANUP_CATEGORIES:
            self.assertTrue(block[key], key)
        self.assertEqual(block["keep_symbols"], "%‰°±")
        # the giveaway bug: "markdown" must not have become a font name
        self.assertEqual(rules["styles"]["title"]["font_eastasia"], "黑体")
        self.assertEqual(rules["styles"]["normal"]["font_eastasia"], "宋体")
        self.assertNotIn("markdown", json.dumps(rules["styles"], ensure_ascii=False))

    def test_no_cleanup_line_means_no_cleanup_block(self):
        rules = textparse.parse_text_rules("标题：黑体 三号\n正文：宋体 小四\n")
        self.assertNotIn("cleanup", rules)

    def test_disable_and_unknown_tokens(self):
        rules = textparse.parse_text_rules("清理：无\n正文：宋体 小四\n")
        self.assertNotIn("cleanup", rules)

        rules = textparse.parse_text_rules("清理：markdown 错别字\n正文：宋体 小四\n")
        self.assertTrue(rules["cleanup"]["markdown"])
        self.assertIn("错别字", rules["extras"]["parse"]["cleanup_unrecognised"])

    def test_keep_symbols_alone_is_not_cleanup(self):
        rules = textparse.parse_text_rules("保留符号：%\n正文：宋体 小四\n")
        self.assertNotIn("cleanup", rules)


def make_dirty_docx(path: str) -> str:
    """A document carrying the kind of markdown residue users actually paste in."""
    document = docx.Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)

    document.add_paragraph("# 2026 年度工作总结报告")
    document.add_paragraph("第一部分 总体情况")
    document.add_paragraph("===")
    document.add_paragraph("本年度整体运行平稳，，各项指标均达到目标。")
    document.add_paragraph("- 完成平台迁移")
    document.add_paragraph("**重点**：收入同比增长 12%")
    document.add_paragraph("■ 装饰性项目符号行")
    document.add_paragraph("零%s宽字符行" % ZWSP)
    document.add_paragraph("")
    document.add_paragraph("")
    document.add_paragraph("")
    document.add_paragraph("结尾段落。")
    document.save(path)
    return path


class TestCleanupEndToEnd(unittest.TestCase):

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="wfa-clean-")
        self.target = make_dirty_docx(os.path.join(self.tmpdir, "dirty.docx"))
        self.presets = os.path.join(self.tmpdir, "presets")
        os.makedirs(self.presets, exist_ok=True)
        with open(os.path.join(self.presets, "清理预设.txt"), "w", encoding="utf-8") as handle:
            handle.write(
                "清理：markdown 不可见字符 重复标点 空白 装饰符号 空段落\n"
                "保留符号：% ‰ ° ±\n"
                "标题：黑体 三号 居中 加粗\n"
                "标题1：黑体 四号 左对齐\n"
                "正文：宋体 小四 1.5倍行距 首行缩进2字符 两端对齐\n")
        with open(os.path.join(self.presets, "干净预设.txt"), "w", encoding="utf-8") as handle:
            handle.write("标题：黑体 三号 居中 加粗\n正文：宋体 小四 1.5倍行距\n")

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _apply(self, preset: str, name: str):
        out = os.path.join(self.tmpdir, name)
        summary = applier.apply_format(
            self.target, "preset:%s" % preset, out,
            report_path=os.path.join(self.tmpdir, name + ".report.json"),
            presets_dir=self.presets)
        with open(summary["report_json"], "r", encoding="utf-8") as handle:
            return summary, json.load(handle), docx.Document(out)

    def test_cleanup_preset_rewrites_text_and_reports_it(self):
        summary, report, after = self._apply("清理预设", "cleaned.docx")
        self.assertTrue(summary["ok"], summary)
        self.assertTrue(summary["text_cleaned"])
        self.assertTrue(summary["text_unchanged"])

        texts = [p.text for p in after.paragraphs]
        self.assertIn("2026 年度工作总结报告", texts)
        self.assertNotIn("# 2026 年度工作总结报告", texts)
        self.assertIn("本年度整体运行平稳，各项指标均达到目标。", texts)
        self.assertIn("完成平台迁移", texts)
        self.assertIn("重点：收入同比增长 12%", texts)
        self.assertIn("装饰性项目符号行", texts)
        self.assertIn("零宽字符行", texts)
        self.assertIn("结尾段落。", texts)
        # "===" was a setext underline: dropped, and it promoted the paragraph
        # directly above it instead of deleting that paragraph's text
        self.assertNotIn("===", texts)
        self.assertIn("第一部分 总体情况", texts)

        # the blank run collapsed to a single empty paragraph
        self.assertEqual(sum(text == "" for text in texts), 1)

        # "第一部分 总体情况" was promoted to heading1 by the setext underline
        promoted = [p for p in after.paragraphs if p.text == "第一部分 总体情况"]
        self.assertEqual(promoted[0].style.name.lower().replace(" ", ""), "heading1")

        block = report["cleanup"]
        self.assertTrue(block["applied"])
        self.assertEqual(block["paragraphs_removed_markers"], 1)
        self.assertEqual(block["paragraphs_removed_blank"], 2)
        self.assertEqual(block["paragraphs_removed"], 3)
        self.assertEqual(block["heading_hints"], 2)
        self.assertGreater(block["paragraphs_changed"], 0)
        self.assertGreater(report["validation"]["chars_removed_by_cleanup"], 0)
        # formatting itself is still verified against the post-cleanup baseline
        self.assertEqual(report["validation"]["text_sha256_baseline"],
                         report["validation"]["text_sha256_after"])
        self.assertNotEqual(report["validation"]["text_sha256_before"],
                            report["validation"]["text_sha256_after"])
        # samples are short and never carry whole paragraphs
        for sample in block["samples"]:
            self.assertLessEqual(len(sample["before"]), applier.SAMPLE_CHARS + 1)

    def test_marker_only_paragraphs_are_dropped(self):
        path = os.path.join(self.tmpdir, "markers.docx")
        document = docx.Document()
        document.add_paragraph("报告标题")
        document.add_paragraph("")
        document.add_paragraph("---")            # horizontal rule (previous line blank)
        document.add_paragraph("| 指标 | 数值 |")
        document.add_paragraph("| --- | --- |")  # table separator row
        document.add_paragraph("```python")      # fence delimiters
        document.add_paragraph("print('a *b* c')")
        document.add_paragraph("```")
        document.add_paragraph("说明文字")
        document.save(path)

        out = os.path.join(self.tmpdir, "markers-out.docx")
        summary = applier.apply_format(
            path, "preset:清理预设", out,
            report_path=os.path.join(self.tmpdir, "markers.report.json"),
            presets_dir=self.presets)
        self.assertTrue(summary["ok"], summary)
        with open(summary["report_json"], "r", encoding="utf-8") as handle:
            report = json.load(handle)

        texts = [p.text for p in docx.Document(out).paragraphs]
        self.assertEqual(texts, ["报告标题", "", "指标 数值", "print('a *b* c')", "说明文字"])
        # HR, table separator and both fence delimiters are gone; the single blank
        # paragraph was already unique, so nothing was collapsed
        self.assertEqual(report["cleanup"]["paragraphs_removed_markers"], 4)
        self.assertEqual(report["cleanup"]["paragraphs_removed_blank"], 0)

    def test_preset_without_cleanup_leaves_text_identical(self):
        summary, report, after = self._apply("干净预设", "untouched.docx")
        self.assertTrue(summary["ok"], summary)
        self.assertFalse(summary["text_cleaned"])
        self.assertIsNone(summary["cleanup"])
        self.assertNotIn("cleanup", report["rules"]["cleanup_categories"])
        self.assertEqual(report["validation"]["text_sha256_before"],
                         report["validation"]["text_sha256_after"])
        self.assertEqual([p.text for p in docx.Document(self.target).paragraphs],
                         [p.text for p in after.paragraphs])

    def test_cleanup_is_visible_through_the_parser(self):
        parsed = parser.parse_format_source("preset:清理预设", presets_dir=self.presets)
        self.assertTrue(parsed["rules"]["cleanup"]["enabled"])
        # a cached parse must round-trip the cleanup block unchanged
        again = parser.parse_format_source("preset:清理预设", presets_dir=self.presets)
        self.assertTrue(again["cache_hit"])
        self.assertEqual(parsed["rules"]["cleanup"], again["rules"]["cleanup"])


class TestSoftReturnParagraphs(unittest.TestCase):
    """Shift+Enter leaves <w:br/>; a paragraph holding only that is still blank."""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="wfa-br-")

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _apply(self, document, text, name="out.docx"):
        target = os.path.join(self.tmpdir, name + ".in.docx")
        document.save(target)
        out = os.path.join(self.tmpdir, name)
        summary = applier.apply_format(
            target, "", out,
            inline_text=text,
            report_path=os.path.join(self.tmpdir, name + ".report.json"))
        self.assertTrue(summary["ok"], summary)
        return summary, docx.Document(out)

    def test_br_only_and_space_only_paragraphs_are_removed(self):
        document = docx.Document()
        document.add_paragraph("第一段正文")
        document.add_paragraph().add_run().add_break()
        document.add_paragraph("第二段正文")
        document.add_paragraph("   ")
        run = document.add_paragraph().add_run()
        run.add_break()
        run.add_break()
        document.add_paragraph("第三段正文")
        wrapped = document.add_paragraph()
        wrapped.add_run("第一行").add_break()
        wrapped.add_run("第二行")

        summary, after = self._apply(
            document, "清理：删除空段落\n正文：宋体 小四 1.5倍行距\n")
        self.assertEqual([p.text for p in after.paragraphs],
                         ["第一段正文", "第二段正文", "第三段正文", "第一行\n第二行"])
        self.assertEqual(summary["cleanup"]["paragraphs_removed"], 3)
        # a soft return inside real text is content and must survive
        self.assertEqual(len(list(after.paragraphs[-1]._p.iter(qn("w:br")))), 1)

    def test_page_break_paragraph_is_protected(self):
        document = docx.Document()
        document.add_paragraph("前")
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_paragraph("后")

        _summary, after = self._apply(
            document, "清理：删除空段落\n正文：宋体 小四\n")
        self.assertEqual(len(after.paragraphs), 3)
        breaks = list(after.paragraphs[1]._p.iter(qn("w:br")))
        self.assertEqual(len(breaks), 1)
        self.assertEqual(breaks[0].get(qn("w:type")), "page")

    def test_spacing_is_zeroed_at_xml_level_and_line_spacing_survives(self):
        document = docx.Document()
        first = document.add_paragraph("正文一")
        spacing = first._p.get_or_add_pPr().get_or_add_spacing()
        spacing.set(qn("w:beforeLines"), "500")
        spacing.set(qn("w:afterLines"), "500")
        spacing.set(qn("w:beforeAutospacing"), "1")
        spacing.set(qn("w:afterAutospacing"), "1")
        document.add_paragraph("正文二")

        _summary, after = self._apply(
            document, "清理：删除空段落\n正文：宋体 小四 1.5倍行距 段前 0 磅 段后 0 磅\n")
        for paragraph in after.paragraphs:
            pPr = paragraph._p.pPr
            spacing = pPr.find(qn("w:spacing"))
            self.assertIsNotNone(spacing)
            self.assertEqual(spacing.get(qn("w:before")), "0")
            self.assertEqual(spacing.get(qn("w:after")), "0")
            # line-based / automatic spacing from the original doc must be gone
            self.assertIsNone(spacing.get(qn("w:beforeLines")))
            self.assertIsNone(spacing.get(qn("w:afterLines")))
            self.assertIsNone(spacing.get(qn("w:beforeAutospacing")))
            self.assertIsNone(spacing.get(qn("w:afterAutospacing")))
            # rebuilding the node must not wipe the 1.5x line spacing
            self.assertEqual(spacing.get(qn("w:line")), "360")
            self.assertEqual(spacing.get(qn("w:lineRule")), "auto")
            ctx = pPr.find(qn("w:contextualSpacing"))
            self.assertIsNotNone(ctx)
            self.assertEqual(ctx.get(qn("w:val")), "0")


class TestFragmentCleanup(unittest.TestCase):
    """Formula fragments must survive cleanup and merge back into one paragraph."""

    def test_equals_and_code_symbols_survive_markdown_cleanup(self):
        clean = cleaner()
        for text in ("C=A×B", "=", "_mm256_*", "__m256",
                     "std::chrono::high_resolution_clock", "串行算法 1.00x \\"):
            self.assertEqual(clean.clean(text)[0], text, text)
        # real markdown emphasis still goes away
        self.assertEqual(clean.clean("这是 **加粗** 与 _斜体_ 测试")[0],
                         "这是 加粗 与 斜体 测试")

    def test_fragment_merge_end_to_end(self):
        tmpdir = tempfile.mkdtemp(prefix="wfa-frag-")
        try:
            target = os.path.join(tmpdir, "fragments.docx")
            document = docx.Document()
            for text in ("实验报告", "一、 问题描述",
                         "矩阵乘法的基础运算为 ", "C", "=", "A", "×", "B",
                         "C=A×B，理论时间复杂度为 ", "O", "(", "n", "3", ")",
                         "O(n ", "3", " )。本实验实现了四种版本：",
                         "后续内容。"):
                document.add_paragraph(text)
            document.save(target)

            presets = os.path.join(tmpdir, "presets")
            os.makedirs(presets, exist_ok=True)
            with open(os.path.join(presets, "碎片预设.txt"), "w", encoding="utf-8") as handle:
                handle.write(
                    "清理：碎片\n"
                    "结构预处理：编号正则\n"
                    "标题1：黑体 四号\n"
                    "正文：宋体 小四\n")

            out = os.path.join(tmpdir, "fragments-out.docx")
            summary = applier.apply_format(
                target, "preset:碎片预设", out,
                report_path=os.path.join(tmpdir, "fragments.report.json"),
                presets_dir=presets)
            self.assertTrue(summary["ok"], summary)
            self.assertTrue(summary["text_unchanged"], summary)

            with open(summary["report_json"], "r", encoding="utf-8") as handle:
                report = json.load(handle)
            block = report["cleanup"]["fragment_merge"]
            self.assertTrue(block["applied"])
            self.assertGreater(block["paragraphs_merged"], 0)
            self.assertGreaterEqual(block["duplicates_dropped"], 5)

            after = docx.Document(out)
            texts = [p.text for p in after.paragraphs]
            merged = ("矩阵乘法的基础运算为 C=A×B，理论时间复杂度为 "
                      "O(n3)。本实验实现了四种版本：")
            self.assertIn(merged, texts)
            for fragment in ("C", "=", "A", "×", "B", "O", "(", "n", "3", ")"):
                self.assertNotIn(fragment, texts)
            self.assertNotIn(" )。本实验实现了四种版本：", texts)
            # the sentence tail must not become a blue Heading 3
            self.assertFalse(any(p.style.name.lower().startswith("heading3")
                                 for p in after.paragraphs))
            merged_paragraph = [p for p in after.paragraphs if p.text == merged][0]
            self.assertEqual(merged_paragraph.style.name, "Normal")
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)


if __name__ == "__main__":
    unittest.main(verbosity=1)
