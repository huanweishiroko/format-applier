"""Tests for 结构预处理 (整段加粗 + 编号正则 -> Word heading styles) and 删除空段落.

The structural pass is the second opt-in pass: it may only change paragraph
styles, never text, so it is tested at three levels - the pure detection helpers,
the parser boundary (``结构预处理：`` lines) and end to end through
``apply_format`` with a real preset.
"""

from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "src")
for path in (SRC, ROOT):
    if path not in sys.path:
        sys.path.insert(0, path)

import docx  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Pt  # noqa: E402

from word_format_applier import applier, cleanup, structure, textparse  # noqa: E402

CONFIG = {"enabled": True, "bold": True, "numbering": True, "ladder": "auto"}


def block(text: str, bold: bool = False):
    """A one-run paragraph body helper; returns a docx Document."""
    document = docx.Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    for item in text.split("\n"):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(item)
        run.font.bold = bold
    return document


def add(document, text: str, bold: bool = False, style: str = None):
    paragraph = document.add_paragraph(style=style) if style else document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.bold = bold
    return paragraph


class TestNumberingMatch(unittest.TestCase):

    def test_official_and_decimal_levels(self):
        cases = {
            "第一章 总则": (1, 1),
            "第一部分 总体情况": (1, 1),
            "一、总体情况": (1, 1),
            "（一）重点工作": (2, 2),
            "(一) 重点工作": (2, 2),
            "1. 平台迁移": (3, 1),
            "1、平台迁移": (3, 1),
            "1.1 迁移范围": (4, 2),
            "1.1.1 迁移步骤": (5, 3),
            "（1）细节": (4, 4),
            "① 说明": (5, 5),
        }
        for text, (official, decimal) in cases.items():
            match = structure.numbering_match(text)
            self.assertIsNotNone(match, text)
            self.assertEqual(match[0], official, text)
            self.assertEqual(match[1], decimal, text)

    def test_prose_is_not_numbering(self):
        for text in ("1.5倍行距说明",
                     "2026 年度工作总结报告",
                     "3.5万元的预算已经批复"):
            self.assertIsNone(structure.numbering_match(text), text)

    def test_compact_dotted_is_bold_only(self):
        # "1.1概述" without a space is only read for a whole-bold paragraph
        self.assertIsNone(structure.numbering_match("1.1概述"))
        match = structure.numbering_match("1.1概述", compact=True)
        self.assertIsNotNone(match)
        self.assertEqual(match[1], 2)


class TestBoldSignal(unittest.TestCase):

    def test_every_non_empty_run_must_be_bold(self):
        document = docx.Document()
        all_bold = document.add_paragraph()
        all_bold.add_run("整段加粗")
        all_bold.add_run("也是粗体").font.bold = True
        all_bold.runs[0].font.bold = True
        self.assertTrue(structure.whole_paragraph_bold(all_bold))

        mixed = document.add_paragraph()
        mixed.add_run("前半").font.bold = True
        mixed.add_run("后半").font.bold = False
        self.assertFalse(structure.whole_paragraph_bold(mixed))

        empty = document.add_paragraph()
        self.assertFalse(structure.whole_paragraph_bold(empty))

    def test_style_level_bold_still_counts(self):
        document = docx.Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("标题走样式加粗")
        document.styles["Normal"].font.bold = True
        self.assertTrue(structure.whole_paragraph_bold(paragraph))


class TestDocumentTitle(unittest.TestCase):
    """首段标题 / 居中短行 -> Word 内置“标题（Title）”样式。"""

    def test_numbered_first_paragraph_is_not_a_title(self):
        document = docx.Document()
        add(document, "一、总体情况")
        add(document, "正文。")
        assignments, meta = structure.detect_levels(
            document.paragraphs, dict(CONFIG, lead_title=True, centered=True))
        self.assertIsNone(meta["title_assigned"])
        self.assertEqual([item["style"] for item in assignments], ["Heading 1"])

    def test_lead_title_beats_the_bold_signal(self):
        document = docx.Document()
        add(document, "2026 年度工作总结报告", bold=True)   # bold AND first
        add(document, "一、总体情况")
        assignments, meta = structure.detect_levels(
            document.paragraphs, dict(CONFIG, lead_title=True))
        self.assertEqual(assignments[0]["style"], "Title")
        self.assertIsNone(assignments[0]["level"])
        self.assertEqual(meta["title_assigned"], "structure:lead_title")
        # the title replaces the bold heading level for that same paragraph
        self.assertEqual([item["style"] for item in assignments[1:]], ["Heading 1"])
        self.assertEqual(meta["candidates"], 2)

    def test_prose_first_paragraph_is_left_alone(self):
        document = docx.Document()
        add(document, "本文档记录了本年度整体运行情况，各项指标均达到既定目标。")
        add(document, "更多正文。")
        assignments, meta = structure.detect_levels(
            document.paragraphs, dict(CONFIG, lead_title=True))
        self.assertEqual(assignments, [])
        self.assertIsNone(meta["title_assigned"])

    def test_centered_line_is_the_fallback(self):
        document = docx.Document()
        add(document, "本年度整体运行平稳，各项核心指标均达到既定目标，收入增长。")
        centered = add(document, "存储芯片市场分析")
        centered.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add(document, "一、总体情况")
        assignments, meta = structure.detect_levels(
            document.paragraphs, dict(CONFIG, lead_title=True, centered=True))
        self.assertEqual(assignments[0]["style"], "Title")
        self.assertEqual(assignments[0]["paragraph"].text, "存储芯片市场分析")
        self.assertEqual(meta["title_assigned"], "structure:centered_title")

    def test_centered_lines_are_only_looked_for_in_the_front_matter(self):
        document = docx.Document()
        add(document, "一、总体情况")
        add(document, "正文。")
        late = add(document, "表 1 关键指标")
        late.alignment = WD_ALIGN_PARAGRAPH.CENTER
        assignments, meta = structure.detect_levels(
            document.paragraphs, dict(CONFIG, centered=True))
        self.assertIsNone(meta["title_assigned"])
        self.assertEqual([item["style"] for item in assignments], ["Heading 1"])

    def test_a_document_that_already_has_a_title_keeps_it(self):
        document = docx.Document()
        add(document, "原有标题", style="Title")
        add(document, "一、总体情况")
        assignments, meta = structure.detect_levels(
            document.paragraphs, dict(CONFIG, lead_title=True, centered=True))
        self.assertTrue(meta["existing_title"])
        self.assertIsNone(meta["title_assigned"])
        self.assertEqual([item["style"] for item in assignments], ["Heading 1"])

    def test_title_signals_are_opt_in(self):
        document = docx.Document()
        add(document, "2026 年度工作总结报告")
        add(document, "一、总体情况")
        assignments, meta = structure.detect_levels(document.paragraphs, CONFIG)
        self.assertIsNone(meta["title_assigned"])
        self.assertEqual([item["style"] for item in assignments], ["Heading 1"])


class TestDetectLevels(unittest.TestCase):

    def test_official_ladder(self):
        document = docx.Document()
        add(document, "一、总体情况")
        add(document, "（一）重点工作")
        add(document, "1. 平台迁移")
        add(document, "这是一段正文，长度正常，以句号结尾。")
        add(document, "注意事项", bold=True)
        add(document, "1. 完成迁移。")            # sentence tail: not a heading
        add(document, "1.5倍行距说明")
        add(document, "3.5 万元的预算已经批复")    # lone "1.1"-shaped prose
        add(document, "一、" + "很长的标题" * 12)  # too long
        assignments, meta = structure.detect_levels(document.paragraphs, CONFIG)
        levels = {item["paragraph"].text: item["level"] for item in assignments}
        self.assertEqual(levels["一、总体情况"], 1)
        self.assertEqual(levels["（一）重点工作"], 2)
        self.assertEqual(levels["1. 平台迁移"], 3)
        self.assertEqual(levels["注意事项"], 1)
        self.assertNotIn("1. 完成迁移。", levels)
        self.assertNotIn("1.5倍行距说明", levels)
        self.assertNotIn("3.5 万元的预算已经批复", levels)
        self.assertEqual(meta["ladder"], "official")
        self.assertEqual(meta["decimal_rejected"], 1)
        self.assertEqual(meta["candidates"], 4)

    def test_decimal_ladder_is_detected_from_the_document(self):
        document = docx.Document()
        add(document, "1. 概述")
        add(document, "1.1 范围")
        add(document, "1.1.1 边界")
        add(document, "1.2 方法")
        add(document, "1.1.2 精度")
        assignments, meta = structure.detect_levels(document.paragraphs, CONFIG)
        levels = {item["paragraph"].text: item["level"] for item in assignments}
        self.assertEqual(meta["ladder"], "decimal")
        self.assertEqual(meta["decimal_rejected"], 0)
        self.assertEqual(levels["1. 概述"], 1)
        self.assertEqual(levels["1.1 范围"], 2)
        self.assertEqual(levels["1.1.1 边界"], 3)
        self.assertEqual(levels["1.2 方法"], 2)

    def test_lone_dotted_number_is_not_an_outline(self):
        document = docx.Document()
        add(document, "3.5 万元的预算已经批复")
        assignments, meta = structure.detect_levels(document.paragraphs, CONFIG)
        self.assertEqual(assignments, [])
        self.assertEqual(meta["ladder"], "official")
        self.assertEqual(meta["decimal_rejected"], 1)

        bold = docx.Document()
        add(bold, "1.1 迁移范围", bold=True)
        assignments, meta = structure.detect_levels(bold.paragraphs, CONFIG)
        self.assertEqual(meta["ladder"], "decimal")
        self.assertEqual(assignments[0]["level"], 2)

    def test_ladder_can_be_pinned(self):
        document = docx.Document()
        add(document, "1. 概述")
        assignments, meta = structure.detect_levels(
            document.paragraphs, dict(CONFIG, ladder="official"))
        self.assertEqual(meta["ladder"], "official")
        self.assertEqual(assignments[0]["level"], 3)

    def test_existing_styles_and_tables_are_left_alone(self):
        document = docx.Document()
        add(document, "一、总体情况", style="Heading 1")
        table = document.add_table(rows=1, cols=1)
        cell_paragraph = table.cell(0, 0).paragraphs[0]
        cell_paragraph.add_run("一、表内编号").font.bold = True
        paragraphs = list(document.paragraphs) + [cell_paragraph]
        assignments, meta = structure.detect_levels(paragraphs, CONFIG)
        self.assertEqual(assignments, [])
        self.assertEqual(meta["skipped_styled"], 1)
        self.assertEqual(meta["skipped_table"], 1)

    def test_globally_bold_document_disables_the_bold_signal(self):
        document = docx.Document()
        for index in range(5):
            add(document, "整篇都是粗体的第 %d 段" % index, bold=True)
        assignments, meta = structure.detect_levels(document.paragraphs, CONFIG)
        self.assertEqual(assignments, [])
        self.assertFalse(meta["bold_signal"])
        self.assertIn("bold_disabled", meta)

    def test_bold_only_config_ignores_numbering(self):
        document = docx.Document()
        add(document, "一、总体情况")
        add(document, "注意事项", bold=True)
        assignments, _ = structure.detect_levels(
            document.paragraphs, {"enabled": True, "bold": True})
        self.assertEqual([item["paragraph"].text for item in assignments], ["注意事项"])

    def test_disabled_config_detects_nothing(self):
        document = docx.Document()
        add(document, "一、总体情况", bold=True)
        assignments, meta = structure.detect_levels(document.paragraphs, {})
        self.assertEqual(assignments, [])
        self.assertFalse(meta["bold_signal"])
        self.assertEqual(meta["candidates"], 0)


class TestStructureDirectives(unittest.TestCase):

    def test_directive_is_not_read_as_a_format_clause(self):
        rules = textparse.parse_text_rules(
            "结构预处理：整段加粗 编号正则\n"
            "标题1：黑体 四号 左对齐\n"
            "正文：宋体 小四 1.5倍行距\n")
        block = rules["structure"]
        self.assertTrue(block["enabled"])
        self.assertTrue(block["bold"])
        self.assertTrue(block["numbering"])
        self.assertEqual(block["ladder"], "auto")
        # "整段加粗" must not have leaked into the paragraph roles
        self.assertNotIn("bold", rules["styles"]["normal"])
        self.assertNotIn("整段加粗", json.dumps(rules["styles"], ensure_ascii=False))

    def test_disable_and_unknown_tokens(self):
        rules = textparse.parse_text_rules("结构预处理：无\n正文：宋体 小四\n")
        self.assertNotIn("structure", rules)

        rules = textparse.parse_text_rules("结构预处理：整段加粗 花里胡哨\n正文：宋体 小四\n")
        self.assertTrue(rules["structure"]["bold"])
        self.assertIn("花里胡哨", rules["extras"]["parse"]["structure_unrecognised"])

    def test_ladder_tokens(self):
        rules = textparse.parse_text_rules("结构预处理：编号正则 小数编号\n正文：宋体 小四\n")
        self.assertEqual(rules["structure"]["ladder"], "decimal")
        self.assertFalse(rules["structure"].get("bold"))

        rules = textparse.parse_text_rules("结构预处理：全部\n正文：宋体 小四\n")
        for key in structure.STRUCTURE_CATEGORIES:
            self.assertTrue(rules["structure"][key], key)

    def test_ladder_alone_is_not_a_pass(self):
        rules = textparse.parse_text_rules("结构预处理：公文编号\n正文：宋体 小四\n")
        self.assertNotIn("structure", rules)


class TestCleanupNumberingInteraction(unittest.TestCase):

    def test_numbering_is_preserved_when_the_reader_asked_for_it(self):
        config = {"enabled": True, "markdown": True}
        kept = cleanup.TextCleaner(config, preserve_numbering=True)
        text, info = kept.clean("1. 平台迁移")
        self.assertEqual(text, "1. 平台迁移")
        self.assertTrue(info.get("numbering"))

        dropped = cleanup.TextCleaner(config)
        self.assertEqual(dropped.clean("1. 平台迁移")[0], "平台迁移")

    def test_bullet_markers_are_still_markdown_residue(self):
        kept = cleanup.TextCleaner({"enabled": True, "markdown": True},
                                   preserve_numbering=True)
        self.assertEqual(kept.clean("- 完成平台迁移")[0], "完成平台迁移")


PRESET = (
    "清理：markdown 删除空段落\n"
    "保留符号：% ‰ ° ±\n"
    "结构预处理：整段加粗 编号正则\n"
    "标题：黑体 三号 居中 加粗\n"
    "标题1：黑体 四号 左对齐\n"
    "标题2：黑体 小四 左对齐\n"
    "标题3：黑体 小四 左对齐\n"
    "正文：宋体 小四 1.5倍行距 首行缩进2字符 两端对齐\n")

#: the same preset with the two document-title signals switched on
TITLE_PRESET = PRESET.replace("结构预处理：整段加粗 编号正则",
                              "结构预处理：整段加粗 编号正则 首段标题 居中短行")

BODY = "本年度整体运行平稳，各项核心指标均达到既定目标，收入同比增长百分之十二。"


def make_document(path: str) -> str:
    document = docx.Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    add(document, "2026 年度工作总结报告", bold=True)
    add(document, "")
    add(document, "")
    add(document, "一、总体情况")
    add(document, BODY)
    add(document, "（一）重点工作", bold=True)
    add(document, BODY)
    add(document, "1. 平台迁移", bold=True)
    add(document, BODY)
    add(document, "")
    add(document, "注意事项", bold=True)
    add(document, "结尾段落。")
    add(document, "")
    document.save(path)
    return path


class TestStructureEndToEnd(unittest.TestCase):

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="wfa-structure-")
        self.target = make_document(os.path.join(self.tmpdir, "target.docx"))
        self.presets = os.path.join(self.tmpdir, "presets")
        os.makedirs(self.presets, exist_ok=True)
        with open(os.path.join(self.presets, "结构预设.txt"), "w", encoding="utf-8") as handle:
            handle.write(PRESET)

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _apply(self, name: str):
        out = os.path.join(self.tmpdir, name)
        summary = applier.apply_format(
            self.target, "preset:结构预设", out,
            report_path=os.path.join(self.tmpdir, name + ".report.json"),
            presets_dir=self.presets)
        with open(summary["report_json"], "r", encoding="utf-8") as handle:
            return summary, json.load(handle), docx.Document(out)

    def test_heading_styles_fonts_and_no_blank_paragraphs(self):
        summary, report, after = self._apply("structured.docx")
        self.assertTrue(summary["ok"], summary)
        self.assertTrue(summary["text_unchanged"])

        by_text = {paragraph.text: paragraph for paragraph in after.paragraphs}
        # every blank paragraph is gone - not just the third of a run
        self.assertEqual([text for text in by_text if text == ""], [])
        self.assertEqual(len(after.paragraphs), 9)

        expected_styles = {
            "2026 年度工作总结报告": "heading1",
            "一、总体情况": "heading1",
            "（一）重点工作": "heading2",
            "1. 平台迁移": "heading3",     # the number survives as document numbering
            "注意事项": "heading1",
        }
        for text, role in expected_styles.items():
            paragraph = by_text[text]
            self.assertEqual(paragraph.style.name.lower().replace(" ", ""), role, text)

        # the ordinary role binding then applies the 黑体 heading fonts
        heading1 = by_text["一、总体情况"].runs[0]
        self.assertEqual(heading1.font.size, Pt(14))          # 四号
        self.assertEqual(heading1._element.rPr.rFonts.get(docx.oxml.ns.qn("w:eastAsia")), "黑体")
        heading2 = by_text["（一）重点工作"].runs[0]
        self.assertEqual(heading2.font.size, Pt(12))          # 小四
        heading3 = by_text["1. 平台迁移"].runs[0]
        self.assertEqual(heading3._element.rPr.rFonts.get(docx.oxml.ns.qn("w:eastAsia")), "黑体")

        # body text keeps the 正文 spec
        body = by_text[BODY]
        self.assertEqual(body.runs[0].font.size, Pt(12))
        self.assertEqual(body.runs[0]._element.rPr.rFonts.get(docx.oxml.ns.qn("w:eastAsia")), "宋体")
        self.assertAlmostEqual(float(body.paragraph_format.line_spacing), 1.5, places=3)

        # report: both passes are visible and the text is verified untouched
        self.assertEqual(report["structure"]["categories"], ["bold", "numbering"])
        self.assertEqual(report["structure"]["ladder"], "official")
        self.assertEqual(report["structure"]["headings_mapped"], 5)
        self.assertEqual(report["structure"]["by_level"], {"1": 3, "2": 1, "3": 1})
        self.assertEqual(report["cleanup"]["paragraphs_removed_blank_all"], 4)
        self.assertEqual(report["rules"]["structure_signals"], ["bold", "numbering"])
        # the structural pass only assigns styles: the post-cleanup baseline is
        # still byte-identical after formatting
        self.assertEqual(report["validation"]["text_sha256_baseline"],
                         report["validation"]["text_sha256_after"])
        self.assertNotEqual(report["validation"]["text_sha256_before"],
                            report["validation"]["text_sha256_after"])

    def test_document_title_becomes_the_word_title_style(self):
        with open(os.path.join(self.presets, "标题预设.txt"), "w", encoding="utf-8") as handle:
            handle.write(TITLE_PRESET)
        out = os.path.join(self.tmpdir, "titled.docx")
        summary = applier.apply_format(
            self.target, "preset:标题预设", out,
            report_path=os.path.join(self.tmpdir, "titled.report.json"),
            presets_dir=self.presets)
        self.assertTrue(summary["ok"], summary)
        self.assertTrue(summary["text_unchanged"])

        after = docx.Document(out)
        title = after.paragraphs[0]
        self.assertEqual(title.text, "2026 年度工作总结报告")
        self.assertEqual(title.style.name.lower().replace(" ", ""), "title")
        run = title.runs[0]
        self.assertEqual(run.font.size, Pt(16))                      # 三号
        self.assertTrue(run.font.bold)
        self.assertEqual(run._element.rPr.rFonts.get(docx.oxml.ns.qn("w:eastAsia")), "黑体")
        self.assertEqual(title.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertAlmostEqual(title.paragraph_format.first_line_indent.pt, 0.0, places=1)

        with open(summary["report_json"], "r", encoding="utf-8") as handle:
            report = json.load(handle)
        self.assertEqual(report["structure"]["title_assigned"], "structure:lead_title")
        self.assertEqual(report["structure"]["title_mapped"], 1)
        self.assertEqual(report["structure"]["headings_mapped"], 4)
        self.assertEqual(report["structure"]["by_level"], {"1": 2, "2": 1, "3": 1})
        self.assertEqual(report["rules"]["structure_signals"],
                         ["bold", "numbering", "lead_title", "centered"])

    def test_preset_without_the_pass_is_unaffected(self):
        with open(os.path.join(self.presets, "干净预设.txt"), "w", encoding="utf-8") as handle:
            handle.write("标题1：黑体 四号 左对齐\n正文：宋体 小四 1.5倍行距\n")
        out = os.path.join(self.tmpdir, "plain.docx")
        summary = applier.apply_format(
            self.target, "preset:干净预设", out,
            report_path=os.path.join(self.tmpdir, "plain.report.json"),
            presets_dir=self.presets)
        self.assertTrue(summary["ok"], summary)
        with open(summary["report_json"], "r", encoding="utf-8") as handle:
            report = json.load(handle)
        self.assertEqual(report["rules"]["structure_signals"], [])
        self.assertFalse(report["structure"]["enabled"])
        self.assertEqual(report["cleanup"], {"enabled": False, "applied": False})
        after = docx.Document(out)
        # nothing removed, nothing restyled: the blank paragraphs are still there
        self.assertEqual(sum(text == "" for text in (p.text for p in after.paragraphs)), 4)
        self.assertEqual(after.paragraphs[3].style.name, "Normal")

    def test_a_document_of_only_blank_paragraphs_keeps_one(self):
        path = os.path.join(self.tmpdir, "blank.docx")
        document = docx.Document()
        add(document, "")
        add(document, "")
        add(document, "")
        document.save(path)
        out = os.path.join(self.tmpdir, "blank-out.docx")
        summary = applier.apply_format(
            path, "preset:结构预设", out,
            report_path=os.path.join(self.tmpdir, "blank.report.json"),
            presets_dir=self.presets)
        self.assertTrue(summary["ok"], summary)
        self.assertEqual([p.text for p in docx.Document(out).paragraphs], [""])

    def test_a_page_break_paragraph_is_never_dropped(self):
        from docx.enum.text import WD_BREAK

        path = os.path.join(self.tmpdir, "breaks.docx")
        document = docx.Document()
        add(document, "第一页正文")
        breaker = document.add_paragraph()
        breaker.add_run().add_break(WD_BREAK.PAGE)
        add(document, "第二页正文")
        add(document, "")
        document.save(path)

        out = os.path.join(self.tmpdir, "breaks-out.docx")
        summary = applier.apply_format(
            path, "preset:结构预设", out,
            report_path=os.path.join(self.tmpdir, "breaks.report.json"),
            presets_dir=self.presets)
        self.assertTrue(summary["ok"], summary)
        after = docx.Document(out)
        self.assertEqual([p.text for p in after.paragraphs], ["第一页正文", "", "第二页正文"])
        self.assertTrue(after.paragraphs[1]._p.findall(
            docx.oxml.ns.qn("w:r") + "/" + docx.oxml.ns.qn("w:br")))


if __name__ == "__main__":
    unittest.main(verbosity=1)
