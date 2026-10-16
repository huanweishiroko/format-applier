"""Test suite for word_format_applier.

Every test builds its own .docx fixtures with python-docx; no network, no LLM
(the default path must stay at zero LLM calls).
"""

from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "src")
for path in (SRC, ROOT):
    if path not in sys.path:
        sys.path.insert(0, path)

import docx  # noqa: E402
from docx.shared import Pt  # noqa: E402

from word_format_applier import applier, parser, presets as presets_mod, textparse  # noqa: E402


def make_target_docx(path: str) -> str:
    document = docx.Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = document.add_paragraph("季度工作报告", style="Title")
    heading = document.add_paragraph("第一部分 概述", style="Heading 1")
    body1 = document.add_paragraph("本季度整体运行平稳，各项指标均达到预期目标。")
    body2 = document.add_paragraph("第二段正文用于验证行距与首行缩进是否被正确套用。")
    quote = document.add_paragraph("引用：数据说明一切。", style="Quote")
    document.add_paragraph("表格说明文字")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "营收"
    table.cell(1, 1).text = "100"

    for paragraph in (body1, body2):
        for run in paragraph.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(11)
    document.save(path)
    return path


def make_template_docx(path: str) -> str:
    """A tiny 'style template': formatted normal/heading1/title paragraphs + margins."""
    document = docx.Document()
    section = document.sections[0]
    from docx.shared import Cm

    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)

    title = document.add_paragraph("样式模板标题", style="Title")
    for run in title.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(22)
        run.font.bold = True
        run._element.rPr.rFonts.set(docx.oxml.ns.qn("w:eastAsia"), "黑体")

    heading = document.add_paragraph("模板一级标题", style="Heading 1")
    for run in heading.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(16)
        run.font.bold = True
        run._element.rPr.rFonts.set(docx.oxml.ns.qn("w:eastAsia"), "黑体")

    normal = document.add_paragraph("模板正文段落，用于提取正文样式。")
    for run in normal.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run._element.rPr.rFonts.set(docx.oxml.ns.qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Pt(24)
    document.save(path)
    return path


def read_reports(summary):
    with open(summary["report_json"], "r", encoding="utf-8") as handle:
        return json.load(handle)


class SampleFixtures:
    """Shared, class scoped fixtures."""

    tmpdir = None

    @classmethod
    def setUpClass(cls):
        cls.tmpdir = tempfile.mkdtemp(prefix="wfa-test-")
        cls.target = make_target_docx(os.path.join(cls.tmpdir, "target.docx"))
        cls.template = make_template_docx(os.path.join(cls.tmpdir, "format.docx"))
        cls.presets_dir = os.path.join(cls.tmpdir, "presets")
        os.makedirs(cls.presets_dir, exist_ok=True)
        with open(os.path.join(cls.presets_dir, "格式要求1.docx"), "wb") as handle:
            with open(cls.template, "rb") as source:
                handle.write(source.read())
        with open(os.path.join(cls.presets_dir, "格式要求3.txt"), "w", encoding="utf-8") as handle:
            handle.write("标题1：黑体 三号 居中\n正文：宋体 小四 1.5倍行距 首行缩进2字符\n")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmpdir, ignore_errors=True)


class TestApplyFromDocxTemplate(SampleFixtures, unittest.TestCase):

    def test_format_applied_and_text_unchanged(self):
        out = os.path.join(self.tmpdir, "out-template.docx")
        summary = applier.apply_format(
            self.target, self.template, out,
            report_path=os.path.join(self.tmpdir, "out-template.report.json"))
        self.assertTrue(summary["ok"], summary)
        self.assertTrue(summary["text_unchanged"])
        self.assertEqual(summary["llm_calls"], 0)
        self.assertGreater(summary["applied"], 0)
        self.assertTrue(os.path.isfile(out))

        before = docx.Document(self.target)
        after = docx.Document(out)
        self.assertEqual(
            "\n".join(p.text for p in before.paragraphs),
            "\n".join(p.text for p in after.paragraphs))
        self.assertEqual(before.tables[0].cell(0, 0).text, after.tables[0].cell(0, 0).text)

        # body paragraph: eastAsia font, size and 1.5 line spacing from the template
        body = after.paragraphs[2]
        run = body.runs[0]
        self.assertEqual(run.font.size, Pt(12))
        self.assertEqual(run._element.rPr.rFonts.get(docx.oxml.ns.qn("w:eastAsia")), "宋体")
        self.assertAlmostEqual(float(body.paragraph_format.line_spacing), 1.5, places=3)
        self.assertAlmostEqual(body.paragraph_format.first_line_indent.pt, 24.0, places=1)

        # heading paragraph got the explicit Heading 1 style and template size
        heading = after.paragraphs[1]
        self.assertEqual(heading.style.name.lower().replace(" ", ""), "heading1")
        self.assertEqual(heading.runs[0].font.size, Pt(16))
        self.assertEqual(heading.runs[0]._element.rPr.rFonts.get(docx.oxml.ns.qn("w:eastAsia")), "黑体")

        # table text follows the body font
        cell_run = after.tables[0].cell(0, 0).paragraphs[0].runs[0]
        self.assertEqual(cell_run.font.size, Pt(12))

        # page setup from the template
        section = after.sections[0]
        self.assertAlmostEqual(section.top_margin.cm, 2.0, places=2)
        self.assertAlmostEqual(section.left_margin.cm, 3.0, places=2)

        report = read_reports(summary)
        self.assertEqual(report["validation"]["text_sha256_before"],
                         report["validation"]["text_sha256_after"])
        self.assertGreater(report["source"]["rule_count"], 0)
        self.assertEqual(report["llm"]["calls"], 0)


class TestTextRequirementParsing(SampleFixtures, unittest.TestCase):

    def test_text_rules_parse(self):
        text = ("标题1：黑体，三号，加粗，居中，段前12磅，段后6磅\n"
                "正文：宋体，小四，行距1.5倍，首行缩进2字符，两端对齐\n"
                "表格文字：宋体，五号\n"
                "一级标题：颜色#FF0000\n")
        rules = textparse.parse_text_rules(text)
        styles = rules["styles"]
        self.assertEqual(styles["heading1"]["font_eastasia"], "黑体")
        self.assertEqual(styles["heading1"]["size_pt"], 16.0)
        self.assertTrue(styles["heading1"]["bold"])
        self.assertEqual(styles["heading1"]["align"], "center")
        self.assertEqual(styles["heading1"]["space_before_pt"], 12.0)
        self.assertEqual(styles["heading1"]["space_after_pt"], 6.0)
        self.assertEqual(styles["normal"]["font_eastasia"], "宋体")
        self.assertEqual(styles["normal"]["size_pt"], 12.0)
        self.assertEqual(styles["normal"]["line_spacing"], 1.5)
        self.assertEqual(styles["normal"]["first_line_indent_chars"], 2.0)
        self.assertEqual(styles["normal"]["align"], "justify")
        self.assertEqual(styles["table_text"]["size_pt"], 10.5)

    def test_text_requirement_end_to_end_zero_llm(self):
        text = "标题1：黑体 三号 加粗 居中\n正文：宋体 小四 1.5倍行距\n"
        out = os.path.join(self.tmpdir, "out-text.docx")
        summary = applier.apply_format(
            self.target, text, out,
            report_path=os.path.join(self.tmpdir, "out-text.report.json"),
            inline_text=text)
        self.assertTrue(summary["ok"], summary)
        self.assertEqual(summary["llm_calls"], 0)
        self.assertTrue(summary["text_unchanged"])
        after = docx.Document(out)
        heading = after.paragraphs[1]
        self.assertEqual(heading.runs[0].font.size, Pt(16))
        self.assertEqual(heading.runs[0]._element.rPr.rFonts.get(docx.oxml.ns.qn("w:eastAsia")), "黑体")
        self.assertAlmostEqual(after.paragraphs[2].runs[0].font.size.pt, 12.0, places=2)


class TestPresets(SampleFixtures, unittest.TestCase):

    def test_preset_list_resolution_and_cache(self):
        names = presets_mod.list_presets(self.presets_dir)
        self.assertIn("格式要求1", names)
        self.assertIn("格式要求3", names)

        entry = presets_mod.resolve_preset("格式要求1", self.presets_dir)
        self.assertTrue(entry["path"].endswith("格式要求1.docx"))
        resolved = presets_mod.resolve_format_source("格式要求1", presets_dir_arg=self.presets_dir)
        self.assertEqual(resolved["kind"], "preset_docx")
        self.assertEqual(resolved["used_preset"], "格式要求1")
        explicit = presets_mod.resolve_format_source("preset:格式要求3", presets_dir_arg=self.presets_dir)
        self.assertEqual(explicit["kind"], "preset_text")
        # an existing path always wins over the preset table
        as_path = presets_mod.resolve_format_source(self.template, presets_dir_arg=self.presets_dir)
        self.assertEqual(as_path["resolved_by"], "path")

    def test_missing_preset_lists_available_names(self):
        with self.assertRaises(presets_mod.PresetError) as ctx:
            presets_mod.resolve_format_source("不存在的预设", presets_dir_arg=self.presets_dir)
        message = str(ctx.exception)
        self.assertIn("格式要求1", message)
        self.assertIn("格式要求3", message)

    def test_preset_parse_cache_avoids_reparse(self):
        first = parser.parse_format_source("格式要求1", presets_dir=self.presets_dir)
        self.assertFalse(first["cache_hit"])
        self.assertEqual(first["llm"]["calls"], 0)

        second = parser.parse_format_source("格式要求1", presets_dir=self.presets_dir)
        self.assertTrue(second["cache_hit"])
        # identical input must produce an identical ruleset, cache hit or not
        self.assertEqual(second["rules"], first["rules"])
        self.assertEqual(second["llm"]["calls"], 0)
        cache_dir = presets_mod.cache_dir(self.presets_dir)
        self.assertTrue(os.path.isdir(cache_dir))
        self.assertTrue([f for f in os.listdir(cache_dir) if f.startswith("rules-")])

    def test_preset_text_file_applies(self):
        out = os.path.join(self.tmpdir, "out-preset-text.docx")
        summary = applier.apply_format(
            self.target, "preset:格式要求3", out,
            report_path=os.path.join(self.tmpdir, "out-preset-text.report.json"),
            presets_dir=self.presets_dir)
        self.assertTrue(summary["ok"], summary)
        self.assertEqual(summary["used_preset"], "格式要求3")
        self.assertEqual(summary["llm_calls"], 0)
        report = read_reports(summary)
        self.assertTrue(report["source"]["preset_sha256"])
        self.assertGreater(report["source"]["rule_count"], 0)


class TestSafety(unittest.TestCase):

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="wfa-safety-")
        self.target = make_target_docx(os.path.join(self.tmpdir, "target.docx"))

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_dry_run_writes_report_only(self):
        out = os.path.join(self.tmpdir, "dry.docx")
        summary = applier.apply_format(self.target, "正文：宋体 小四", out,
                                       dry_run=True, inline_text="正文：宋体 小四")
        self.assertFalse(os.path.exists(out))
        self.assertTrue(os.path.isfile(summary["report_json"]))
        report = read_reports(summary)
        self.assertTrue(report["dry_run"])
        self.assertTrue(report["validation"]["text_unchanged"])

    def test_refuses_to_overwrite_target(self):
        with self.assertRaises(applier.FormatError):
            applier.apply_format(self.target, "正文：宋体 小四",
                                 self.target, inline_text="正文：宋体 小四")

    def test_llm_fallback_is_off_by_default(self):
        summary = applier.apply_format(
            self.target, "把文档弄得好看一点", os.path.join(self.tmpdir, "weird.docx"),
            report_path=os.path.join(self.tmpdir, "weird.report.json"),
            inline_text="把文档弄得好看一点")
        self.assertEqual(summary["llm_calls"], 0)
        report = read_reports(summary)
        self.assertEqual(report["llm"]["calls"], 0)
        self.assertIn(report["llm"]["reason"], ("disabled", "not_needed:coverage_ok:0.000"))


class TestTextFingerprint(unittest.TestCase):

    def test_fingerprint_ignores_formatting(self):
        tmpdir = tempfile.mkdtemp(prefix="wfa-fp-")
        try:
            path = make_target_docx(os.path.join(tmpdir, "t.docx"))
            document = docx.Document(path)
            before = applier.docxio.text_fingerprint(document)
            for paragraph in document.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(30)
                    run.font.bold = True
            after = applier.docxio.text_fingerprint(document)
            self.assertEqual(before["sha256"], after["sha256"])
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)


if __name__ == "__main__":
    unittest.main(verbosity=1)
