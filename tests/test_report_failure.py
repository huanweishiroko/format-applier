"""回归测试：report.json 失败绝不能拖垮主流程，失败原因必须能被界面看到。

对应事故：
1. docx 已写不出（低完整性/沙箱/杀软）时 CLI 只回 ``ok=false``，stderr 全空，
   界面只能弹「处理失败（退出码 1）」；
2. 报告移动失败被当成致命错误。
"""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from unittest import mock

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "src")
for path in (SRC, ROOT):
    if path not in sys.path:
        sys.path.insert(0, path)

import docx  # noqa: E402

from word_format_applier import applier, cli, report as report_mod  # noqa: E402

FORMAT_TEXT = "正文：宋体 小四 1.5倍行距 首行缩进2字符\n"


def make_docx(path: str) -> str:
    document = docx.Document()
    document.add_paragraph("季度工作报告", style="Title")
    document.add_paragraph("第一部分 概述", style="Heading 1")
    document.add_paragraph("本季度整体运行平稳。")
    document.save(path)
    return path


class ReportFailureTest(unittest.TestCase):
    def setUp(self) -> None:
        self.tmpdir = tempfile.mkdtemp(prefix="wfa-report-fail-")
        self.addCleanup(lambda: __import__("shutil").rmtree(self.tmpdir, ignore_errors=True))
        self.target = make_docx(os.path.join(self.tmpdir, "target.docx"))

    def test_report_write_failure_keeps_document_ok(self) -> None:
        out = os.path.join(self.tmpdir, "out.docx")
        failure = PermissionError(13, "Permission denied", "out.docx.report.json")
        with mock.patch.object(report_mod, "write_reports", side_effect=failure):
            summary = applier.apply_format(
                self.target, "", out,
                inline_text=FORMAT_TEXT,
                report_path=os.path.join(self.tmpdir, "out.docx.report.json"),
                logger=None)
        self.assertTrue(summary["ok"], summary)
        self.assertTrue(os.path.isfile(out), "docx 必须已经落盘")
        self.assertFalse(summary["partial"], summary)
        self.assertIsNone(summary["report_json"], "报告没写出来就不能给假路径")
        self.assertIn("PermissionError", summary["report_error"])
        self.assertTrue(
            any("报告文件移动失败" in item for item in summary["warnings"]),
            summary["warnings"])


class CliFailureMessageTest(unittest.TestCase):
    def setUp(self) -> None:
        self.tmpdir = tempfile.mkdtemp(prefix="wfa-cli-msg-")
        self.addCleanup(lambda: __import__("shutil").rmtree(self.tmpdir, ignore_errors=True))
        self.target = make_docx(os.path.join(self.tmpdir, "target.docx"))
        self.output = os.path.join(self.tmpdir, "out.docx")

    def _run(self, summary: dict):
        argv = ["--target", self.target, "--format", FORMAT_TEXT,
                "--output", self.output, "--json", "--no-write-probe"]
        out, err = io.StringIO(), io.StringIO()
        with mock.patch.object(applier, "apply_format", return_value=summary):
            with redirect_stdout(out), redirect_stderr(err):
                code = cli.main(argv)
        return code, out.getvalue(), err.getvalue()

    def _summary(self, **overrides) -> dict:
        summary = {"ok": True, "output_docx": self.output, "report_json": None,
                   "applied": 1, "skipped": 0, "llm_calls": 0,
                   "text_unchanged": True, "warnings": [], "errors": [],
                   "partial": False}
        summary.update(overrides)
        return summary

    def test_report_only_failure_exits_zero_with_warn(self) -> None:
        summary = self._summary(report_error="PermissionError: [Errno 13] report.json")
        code, out, err = self._run(summary)
        self.assertEqual(code, 0, err)
        self.assertIn("[WARN] 报告文件移动失败，已跳过", err)
        self.assertTrue(json.loads(out)["ok"])

    def test_write_failure_surfaces_real_error_on_stderr(self) -> None:
        summary = self._summary(
            ok=False, partial=True, output_docx=None,
            errors=["failed to write output: C:\\out\\out.docx（PermissionError: [Errno 13] "
                    "Permission denied: 'C:\\\\out\\\\out.docx'）"])
        code, _out, err = self._run(summary)
        self.assertEqual(code, 1)
        self.assertIn("error:", err)
        self.assertIn("PermissionError", err)
        self.assertIn("failed to write output", err)


class WriteProbeTest(unittest.TestCase):
    def test_writable_dir(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            probe = cli.probe_output_dir(os.path.join(tmp, "sub", "out.docx"))
            self.assertTrue(probe["ok"], probe)
            leftovers = [name for _root, _dirs, files in os.walk(tmp)
                         for name in files if "write-probe" in name]
            self.assertEqual(leftovers, [], "探测文件必须删干净")

    def test_unwritable_dir(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            blocker = os.path.join(tmp, "blocker")
            with open(blocker, "w", encoding="utf-8") as handle:
                handle.write("x")
            probe = cli.probe_output_dir(os.path.join(blocker, "out.docx"))
            self.assertFalse(probe["ok"], probe)
            self.assertTrue(probe["error"])
            self.assertIn("不可写", cli.unwritable_message(probe))

    def test_probe_write_cli_exit_codes(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            out = io.StringIO()
            with redirect_stdout(out), redirect_stderr(io.StringIO()):
                code = cli.main(["--probe-write", tmp])
            self.assertEqual(code, 0)
            self.assertTrue(json.loads(out.getvalue())["ok"])


if __name__ == "__main__":
    unittest.main()
